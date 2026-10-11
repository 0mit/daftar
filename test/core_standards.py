#!/usr/bin/env python3
"""The core's law is its own source (v1 part 13): what the files of core/law/ promise, held to them.

Until v1.0.0 the standards' tables were generated from today's law, and this suite held them equal to it. std-vocab is
gone; the files are the law, and this holds what they promise of themselves:
- no file of core/law/ says it is generated, and none names today's law as its source;
- every table of rows has its form — the standards' files under `forms`, the files of a line, a measure and the flow law
  under `columns` — and every row holds its form's required columns and no other: a stray column, or a required one
  missing, is refused by the law's own proof (core/check.py --law);
- the core's names of tables are the ones the core reads (`systems`, `protocols`, `types`, `gaps`), and no form names a
  table by today's name;
- a value the core renamed (soma, lekton, derived, said) is written in the core's word, and prose as written;
- the units: every unit measures a quantity of quantities.yaml, by its English name, with a factor of two whole numbers;
- the standards are read from core/law/ alone: a release's systems, protocols, quantities, currencies, zones and schemes,
  and its gate reads them.
"""
import json, os, re, shutil, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from core import read, standards  # noqa: E402
from core.law import columns_problems  # noqa: E402

FAILS = []
PY = sys.executable
LAW_DIR = os.path.join(ROOT, 'core', 'law')


def check(name, cond, detail=""):
    print(("PASS  " if cond else "FAIL  ") + name + ("" if cond else "  — " + str(detail)[:1500]))
    if not cond:
        FAILS.append(name)


def text(name):
    with open(os.path.join(LAW_DIR, name), encoding='utf-8') as fh:
        return fh.read()


FILES = sorted(f for f in os.listdir(LAW_DIR) if f.endswith('.yaml'))

# ------------------------------------------------------------------------------------------- the source
said = [f for f in FILES if 'GENERATED' in text(f) or 'std-vocab' in text(f)]
check(f"no file of core/law/ ({len(FILES)}) says it is generated, or names today's law as its source", not said, said)
check("...and no generator is left: core/translate.py writes a garden's copy, never a file of the law",
      not re.search(r"(?m)^def (generated|law_as_strings|old_law)\(", open(os.path.join(ROOT, 'core', 'translate.py'),
                                                                        encoding='utf-8').read()))

# ------------------------------------------------------------------------------------------- each table in its form
found = columns_problems(LAW_DIR)
n_tables = sum(1 for f in FILES for k, v in read.data(os.path.join(LAW_DIR, f)).items()
               if f in ('systems.yaml', 'places.yaml', 'protocols.yaml', 'quantities.yaml', 'registries.yaml', 'lines.yaml',
                        'measures.yaml', 'flows.yaml') and isinstance(v, list) and v and all(isinstance(r, dict) for r in v))
check(f"every one of the {n_tables} tables of rows has its form, and every row holds the required columns and no other",
      n_tables >= 25 and not found, found[:6])

tmp = tempfile.mkdtemp(prefix='core-standards-')
try:
    rel = os.path.join(tmp, 'release')
    shutil.copytree(ROOT, rel, ignore=shutil.ignore_patterns('.git', '__pycache__', 'site'))
    r = subprocess.run([PY, os.path.join(rel, 'core', 'check.py'), '--law'], capture_output=True, text=True,
                       encoding='utf-8', cwd=rel)
    check("the core's law is whole in a release's copy: core/check.py --law, 0 errors", r.returncode == 0,
          (r.stdout + r.stderr)[-800:])
    for name, old, new, says in (
            ('systems.yaml', '    neighbours: none\n', '    neighbours: none\n    remark: x\n', "column `remark` is not in"),
            ('flows.yaml', '  - { method: ', '  - { stray: y, method: ', "column `stray` is not in"),
            ('lines.yaml', '  - { aggregate: ', '  - { meaningless: z, aggregate: ', "column `meaningless` is not in")):
        p = os.path.join(rel, 'core', 'law', name)
        was = open(p, encoding='utf-8').read()
        with open(p, 'w', encoding='utf-8', newline='\n') as fh:
            fh.write(was.replace(old, new, 1))
        r = subprocess.run([PY, os.path.join(rel, 'core', 'check.py'), '--law'], capture_output=True, text=True,
                           encoding='utf-8', cwd=rel)
        check(f"...a row of {name} holding a column its form lacks is refused by the law's proof, naming it",
              r.returncode == 1 and says in r.stdout, r.stdout[-600:])
        with open(p, 'w', encoding='utf-8', newline='\n') as fh:
            fh.write(was)
    p = os.path.join(rel, 'core', 'law', 'measures.yaml')
    was = open(p, encoding='utf-8').read()
    with open(p, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write(re.sub(r'(?m)^(  - \{ placement: place), takes: none \}', r'\1 }', was, count=1))
    r = subprocess.run([PY, os.path.join(rel, 'core', 'check.py'), '--law'], capture_output=True, text=True,
                       encoding='utf-8', cwd=rel)
    check("...and a row missing a required column is refused, naming it", r.returncode == 1 and "holds no `takes`" in r.stdout,
          r.stdout[-600:])
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ------------------------------------------------------------------------------------------- the core's names
texts = ''.join(text(f) for f in FILES)
old_names = [w for w in ('anchor_systems', 'net_protocols', 'value_types', 'gap_tokens') if f"registry: {w}" in texts]
check("no form in core/law/ names a table by today's name (anchor_systems, net_protocols, value_types, gap_tokens)",
      not old_names, old_names)
L = read.data(os.path.join(LAW_DIR, 'lines.yaml'))
check("the forms a value is written in and a series' gaps are the core's tables, `types` and `gaps`, each with its form",
      {'count', 'kebab', 'text', 'rows'} <= {r.get('type') for r in L.get('types') or []}
      and {'-', '?', '_'} <= {r.get('token') for r in L.get('gaps') or []}
      and set(L.get('columns') or {}) >= {'types', 'gaps'})
shape = read.data(os.path.join(LAW_DIR, 'systems.yaml'))['system_shape']['sources']
check("a value the core renamed is written in its word: the shape's sources are acts `derive`, `read`, `say` of a "
      "`body` or a `sayable`, never derived/said/soma/lekton",
      shape['reckoning']['arithmetic'] == {'act': 'derive', 'nature': 'sayable'}
      and shape['reckoning']['astronomical'] == {'act': 'read', 'nature': 'body'}
      and shape['crosswalk']['table'] == {'act': 'say', 'nature': 'sayable'}
      and shape['datum']['host']['nature'] == ['body', 'sayable'], shape)
sysrow = standards.here().systems['unix-filesystem']
check("...and prose is kept as written, every value a string: `establishes` is the string `false`, a pattern its own "
      "backslashes", sysrow['establishes'] == 'false'
      and standards.here().systems['git-remote']['pattern'].endswith(r'[^ ]*\.git)|[a-z][a-z0-9+.-]*://[^ ]+)$'),
      sysrow)

# ------------------------------------------------------------------------------------------- the units
quantities = {q['quantity'] for q in read.data(os.path.join(LAW_DIR, 'quantities.yaml'))['quantities']}
units = read.data(os.path.join(LAW_DIR, 'units.yaml'))['units']
bad = [u['unit'] for u in units if u.get('quantity') not in quantities or not u.get('name')
       or not ((isinstance(u.get('factor'), list) and len(u['factor']) == 2 and all(str(x).isdigit() for x in u['factor']))
               or ('factor' not in u and u.get('why')))]
check(f"every one of the core's {len(units)} units in UCUM measures a quantity of quantities.yaml, by its English name, "
      f"with a factor of two whole numbers — or none, saying why (a radian meets a degree only through π; an IU, a "
      f"referenced level and a count of one thing convert into nothing else)", len(units) >= 55 and not bad
      and sum(1 for u in units if 'factor' not in u) == 14, bad)

# ------------------------------------------------------------------------------------------- read from core/law/ alone
S = standards.here()
check(f"the standards read from core/law/: {len(S.systems)} systems, {len(S.protocols)} protocols, {len(S.quantities)} "
      f"quantities, {len(S.currencies)} currencies, {len(S.zones)} zones, {len(S.knowledge.schemes)} schemes",
      len(S.systems) > 50 and len(S.protocols) > 20 and len(S.quantities) > 15 and 'EUR' in S.currencies
      and 'Asia/Tehran' in S.zones and {'isco-08', 'isced-f-2013'} <= set(S.knowledge.schemes)
      and not S.code('isco-08:2522'),
      (len(S.systems), len(S.protocols), len(S.quantities), len(S.currencies), len(S.zones), S.code('isco-08:2522')))
check("...and the release carries no today's law: seed/std-vocab.md is gone",
      not os.path.exists(os.path.join(ROOT, 'seed', 'std-vocab.md')))

# ------------------------------------------------------------------------------------------- a table fetched (Y7)
sys.path.insert(0, os.path.join(ROOT, 'bin'))
import fetch as dmfetch  # noqa: E402
import knowledge as dmknowledge  # noqa: E402
reg = read.data(os.path.join(LAW_DIR, 'registries.yaml'))
isco = next(r for r in reg['knowledge_schemes'] if r['scheme'] == 'isco-08')
check("ISCO-08 is held fetched: from the ILO's own download, read by a reader the law's shape lists and this release has",
      isco['holding'] == 'fetched' and isco['fetch']['from'].startswith('https://webapps.ilo.org/')
      and isco['fetch']['reader'] in reg['fetch_shape']['readers'] and isco['fetch']['reader'] in dmfetch.READERS
      and set(reg['fetch_shape']['readers']) <= set(dmfetch.READERS), isco.get('fetch'))
CACHE = tempfile.mkdtemp(prefix='core-standards-cache-')
os.environ['DAFTAR_STANDARDS_CACHE'] = CACHE
check("...and a machine that fetched none reads the release's copy meanwhile",
      not dmfetch.fetched('isco-08') and (dmknowledge.Knowledge(ROOT).row('isco-08', '0110') or {}).get('name_en')
      == 'Commissioned Armed Forces Officers', dmknowledge.Knowledge(ROOT).row('isco-08', '0110'))
import io  # noqa: E402
import openpyxl  # noqa: E402 — the provider's file is a spreadsheet; openpyxl writes one as the ILO's is written


def xlsx(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


ilo = xlsx([('Level', 'ISCO 08 Code', 'Title EN', 'Definition'), (1, '0', 'Armed Forces Occupations', 'x'),
            (2, '01', 'Commissioned Armed Forces Officers', 'x'), (3, '011', 'Commissioned Armed Forces Officers', 'x'),
            (4, '0110', 'Commissioned Armed Forces Officers, as the provider titles it', 'x')])
src = os.path.join(CACHE, 'ilo.xlsx')
with open(src, 'wb') as fh:
    fh.write(ilo)
r = subprocess.run([sys.executable, os.path.join(ROOT, 'bin', 'fetch.py'), 'isco-08', '--from', src], capture_output=True,
                   text=True, encoding='utf-8', env=dict(os.environ))
got = dmfetch.fetched('isco-08')
import hashlib  # noqa: E402
check("fetch: the provider's file is kept unchanged under its SHA-256, read into the rows the scheme's file holds, and the "
      "line for the save's entry printed",
      r.returncode == 0 and got and got[0][1] == hashlib.sha256(ilo).hexdigest()
      and open(got[0][2], encoding='utf-8').readline().strip().split('\t') == ['code', 'level', 'parent', 'name_en']
      and f"sha-256 {hashlib.sha256(ilo).hexdigest()}" in r.stdout and '- read: isco-08' in r.stdout,
      r.stdout + r.stderr)
check("...and the codes are read from this machine's copy, the provider's titles with them",
      (dmknowledge.Knowledge(ROOT).row('isco-08', '0110') or {}).get('name_en')
      == 'Commissioned Armed Forces Officers, as the provider titles it'
      and (dmknowledge.Knowledge(ROOT).row('isco-08', '01') or {}).get('level') == 'sub-major',
      dmknowledge.Knowledge(ROOT).row('isco-08', '0110'))
for name, data, says in (("a file whose columns are not the ones read", xlsx([('Code', 'Title'), ('0', 'x')]), 'columns'),
                         ("a row whose code is not its level's", xlsx([('Level', 'ISCO 08 Code', 'Title EN'),
                                                                        (2, '0', 'x')]), 'no group read here')):
    try:
        dmfetch.keep('isco-08', data, 'test')
        refused = ''
    except dmfetch.Refused as e:
        refused = str(e)
    check(f"fetch: {name} is refused, and nothing is kept", says in refused and len(dmfetch.fetched('isco-08')) == 1,
          refused)
try:
    dmfetch.keep('isced-f-2013', b'x', 'test')
    refused = ''
except dmfetch.Refused as e:
    refused = str(e)
check("fetch: a scheme the law holds shipped is not fetched", 'not fetched' in refused, refused)
shutil.rmtree(CACHE, ignore_errors=True)

# ------------------------------------------------------------------------------------------- a standard's units crossed
# A SCHEME'S CODES CROSS INTO THE LAW'S UNITS AND KINDS (registries.yaml `units_crosswalk_form`): its row names a
# registry of rows, one a code, each crossing to one target of the law or, with none, saying why. A count written with a
# code is read as the law writes it, and written back with the code that crosses to it (bin/units.py).
import units as dmunits  # noqa: E402
from core.check import garden_law  # noqa: E402
from core import frame  # noqa: E402
from core.measures import Measures, quantity_why  # noqa: E402
sys.path.insert(0, os.path.join(ROOT, 'test'))
import grow  # noqa: E402 — a release of the core made from this tree, and a garden grown from it
LAWNOW = garden_law(ROOT)
crossed = {s: dmunits.crossing(s) for s, r in dmknowledge.Knowledge(ROOT).schemes.items() if r.get('units')}
check(f"the law crosses {len(crossed)} scheme's codes into its units and kinds, each by a registry the law declares, in "
      f"`units_crosswalk_form`'s columns ({', '.join(sorted(crossed))})",
      'unece-rec20' in crossed and not LAWNOW.units_crosswalk_problems()
      and all(set(k for r in rows.values() for k in r) <= set(reg['units_crosswalk_form']) for rows in crossed.values()),
      LAWNOW.units_crosswalk_problems()[:4])
M = Measures(LAWNOW)
trips, judged = [], []
for scheme, rows in sorted(crossed.items()):
    for (axis, code), row in sorted(rows.items()):
        t = dmunits.target(row)
        if axis or not t or not ({'unit', 'in'} & set(t)):
            continue
        count = '3' if 'level' in t or t.get('in') == 'ordinal-number' else '21.5'
        sample = code[:-1] + '2' if row.get('decimals') == 'n' else code
        got = dmunits.read(scheme, '002150' if row.get('decimals') == 'n' else count, sample)
        back = dmunits.write(scheme, got)
        want = dmunits.code_of(scheme, t)[0]
        want = want[:-1] + '2' if row.get('decimals') == 'n' else want
        if back[1] != want or (row.get('decimals') != 'n' and back[0] != count):
            trips.append(f"{scheme}:{code} {count} -> {got} -> {back}")
        why = quantity_why(LAWNOW, got) if 'unit' in t else (
            M.extent(code, got) if 'level' in t else ([] if frame.read(got, LAWNOW.systems).system == t['in'] else ['?']))
        if why:
            judged.append(f"{scheme}:{code} -> {got}: {why}")
check("every code crossed to a unit, a system's positions or its cells is read as the law writes it and written back as "
      "its code — or, of several crossing to one target, the one marked `back`",
      not trips and len([1 for rows in crossed.values() for r in rows.values() if dmunits.target(r)]) >= 150, trips[:6])
check("...and what it is read as is judged as the gate judges it: a quantity in its form, an extent of cells "
      "(`extent_form`), a position of the system the code is the unit of", not judged, judged[:6])

# the form's refusals, in a release's copy: each breach named by the law's proof
tmp = tempfile.mkdtemp(prefix='core-standards-crossing-')
try:
    rel = os.path.join(tmp, 'release')
    shutil.copytree(ROOT, rel, ignore=shutil.ignore_patterns('.git', '__pycache__', 'site'))
    p = os.path.join(rel, 'seed', 'knowledge', 'units-unece-rec20.tsv')
    good = open(p, encoding='utf-8').read()
    for name, rows, says in (
            ("a unit the law has not", "XX1\tkg2\t\t\t\t\t\n", "'kg2' is no unit of the law"),
            ("two targets", "XX1\tkg\tcelsius-scale\t\t\t\t\n", "crosses to unit and in: one target a row"),
            ("a level its system has not", "XX1\t\tgregorian-civil\tfortnight\t\t\t\n", "'fortnight' is no level of gregorian-civil"),
            ("a system that writes no position from a count, named alone", "XX1\t\tgregorian-civil\t\t\t\t\n",
             "gregorian-civil writes no position from a count"),
            ("a row crossing to nothing, saying no why", "XX1\t\t\t\t\t\t\n", "crosses to nothing, and says no `why`"),
            ("`vacant` on a row with a target", "XX1\tkg\t\t\t\tuniversal\tx\n", "on a row that crosses to nothing"),
            ("a code written twice", "KGM\tg\t\t\t\t\t\n", "is written twice"),
            ("a second code crossing to a target, and none marked `back`", "XX1\tkg\t\t\t\t\t\n",
             "KGM, XX1 cross to unit: kg, and none is marked `back`"),
            ("`back` on the one code crossing to its target", "XX1\t{frame}\t\t\ttrue\t\t\n",
             "XX1: is marked `back`, and is the one code crossing to unit: {frame}"),
            ("two codes crossing to one target, one marked `back`", "XX1\tWb\t\t\ttrue\t\t\n", "")):
        with open(p, 'w', encoding='utf-8', newline='\n') as fh:
            fh.write(good + rows)
        r = subprocess.run([PY, os.path.join(rel, 'core', 'check.py'), '--law'], capture_output=True, text=True,
                           encoding='utf-8', cwd=rel)
        if not says:     # WEB crosses to Wb already: a second, marked `back`, is the one of two written back
            check(f"crossing: {name} hold together", r.returncode == 0, r.stdout[-400:])
            continue
        check(f"crossing: {name} is refused by the law's proof, naming it", r.returncode == 1 and says in r.stdout,
              r.stdout[-500:])
    with open(p, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write(good.replace('code\tunit\tin', 'code\tunit\tfactor\tin', 1))
    r = subprocess.run([PY, os.path.join(rel, 'core', 'check.py'), '--law'], capture_output=True, text=True,
                       encoding='utf-8', cwd=rel)
    check("crossing: a column the form has not is refused, naming it (a factor is the unit's, never the crossing's)",
          r.returncode == 1 and "factor: no column of `units_crosswalk_form`" in r.stdout, r.stdout[-500:])
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# UNECE Recommendation 20: trade's units, held fetched from UN/CEFACT, and crossed by the law
rec = next(r for r in reg['knowledge_schemes'] if r['scheme'] == 'unece-rec20')
check("UNECE Recommendation 20 is held fetched from UN/CEFACT's own spreadsheet, by a reader the law lists and this "
      "release has, and the release carries no copy of it: its registry names no file",
      rec['holding'] == 'fetched' and rec['fetch']['from'].startswith('https://unece.org/')
      and rec['fetch']['reader'] == 'unece-rec20-xlsx' and rec['fetch']['reader'] in dmfetch.READERS
      and rec['fetch']['reader'] in reg['fetch_shape']['readers'] and rec['units'] == 'units-unece-rec20'
      and not next(r for r in reg['registry_files'] if r['registry'] == 'unece-rec20').get('file'), rec)
CACHE = tempfile.mkdtemp(prefix='core-standards-cache-')
os.environ['DAFTAR_STANDARDS_CACHE'] = CACHE
try:
    dmknowledge.Knowledge(ROOT).row('unece-rec20', 'KGM')
    unread = ''
except KeyError as e:
    unread = str(e)
r = subprocess.run([PY, os.path.join(ROOT, 'bin', 'knowledge.py'), 'show', 'unece-rec20', 'KGM'], capture_output=True,
                   text=True, encoding='utf-8', env=dict(os.environ))
check("...a machine that fetched none reads none of its codes, and says how to fetch them — the tool, too, without a "
      "traceback", 'has fetched none: python3 bin/fetch.py unece-rec20' in unread and r.returncode == 1
      and 'Traceback' not in r.stdout + r.stderr and 'fetched none' in r.stdout, (unread, r.stdout + r.stderr))
check("...and its crossing is the law's own, shipped: a count written in its codes is read without it",
      dmunits.read('unece-rec20', '12.5', 'KGM') == {'count': '12.5', 'unit': 'kg'})


def rec20(annex2, annex1=()):
    wb = openpyxl.Workbook()
    wb.active.title = 'Intro'
    wb.active.append(('CODES FOR UNITS OF MEASURE USED IN INTERNATIONAL TRADE',))
    a1 = wb.create_sheet('Annex I')
    a1.append(('Group Number', 'Sector', 'Group ID', 'Quantity', 'Level/ Category', 'Status', 'Common Code', 'Name',
               'Conversion Factor', 'Symbol', 'Description'))
    for row in annex1:
        a1.append(row)
    a2 = wb.create_sheet('Annex II & Annex III')
    a2.append(('Status', 'Common\nCode', 'Name', 'Description', 'Level /_x000D_\nCategory', 'Symbol', 'Conversion Factor'))
    for row in annex2:
        a2.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


A2 = [(None, 'KGM', 'kilogram', 'A unit of mass equal to one thousand grams.', '1', 'kg', 'kg'),
      (None, 'MON', 'month', 'Unit of time equal to 1/12 of a year of 365,25 days.', '2', 'mo', '2,629 800 x 10⁶ s'),
      (None, 'C62', 'one', 'Synonym: unit', '1', '1', '1'), ('D', 'HAR', 'hectare', None, '2', 'ha', '10⁴ m²'),
      ('X', '05', 'lift', None, '3.9', None, None), (None, 'XAE', 'aerosol', None, '3.9', None, None)]
A1 = [('01', 'Space and Time', '1', 'mass', '1', None, 'KGM', 'kilogram', 'kg', 'kg', None),
      ('04', 'Mechanics', '9', 'effective mass', '1', None, 'KGM', 'kilogram', 'kg', 'kg', None)]
src = os.path.join(CACHE, 'rec20.xlsx')
with open(src, 'wb') as fh:
    fh.write(rec20(A2, A1))
r = subprocess.run([PY, os.path.join(ROOT, 'bin', 'fetch.py'), 'unece-rec20', '--from', src], capture_output=True,
                   text=True, encoding='utf-8', env=dict(os.environ))
got = dmfetch.fetched('unece-rec20')
check("fetch: UNECE's file is kept unchanged, its sheet of every code read — not Annex I, which lists a code once for "
      "each quantity — into code, status, name and symbol, a deleted and a deprecated code so marked",
      r.returncode == 0 and got and open(got[0][2], encoding='utf-8').readline().strip().split('\t')
      == ['code', 'status', 'name_en', 'symbol'] and '6 rows' in r.stdout
      and (dmknowledge.Knowledge(ROOT).row('unece-rec20', 'HAR') or {}).get('status') == 'deprecated'
      and (dmknowledge.Knowledge(ROOT).row('unece-rec20', '05') or {}).get('status') == 'deleted'
      and (dmknowledge.Knowledge(ROOT).row('unece-rec20', 'KGM') or {}).get('name_en') == 'kilogram', r.stdout + r.stderr)
check("...and it says what the copy holds of the law's crossing: how many of its codes cross, each crossed code it has "
      "not, and each it marks deprecated",
      "4 of its 6 codes cross into the law's units and kinds" in r.stdout and 'not in this copy: ' in r.stdout
      and 'KWH' in r.stdout.split('not in this copy: ')[1].split('\n')[0]
      and 'marked by the provider: HAR (deprecated)' in r.stdout, r.stdout)
for name, data, says in (("a file with no sheet of every code", xlsx([('Code', 'Title')]), 'list every code once'),
                         ("a row whose code is no code", rec20([(None, 'kilo gram', 'x', None, '1', None, None)]),
                          'no code read here'),
                         ("a code written twice", rec20([A2[0], A2[0]]), 'a second row of it')):
    try:
        dmfetch.keep('unece-rec20', data, 'test')
        refused = ''
    except dmfetch.Refused as e:
        refused = str(e)
    check(f"fetch: UNECE's file — {name} — is refused, and nothing is kept",
          says in refused and len(dmfetch.fetched('unece-rec20')) == 1, refused)
r = subprocess.run([PY, os.path.join(ROOT, 'bin', 'fetch.py'), 'unece-rec20'], capture_output=True, text=True,
                   encoding='utf-8', env=dict(os.environ, http_proxy='http://127.0.0.1:9', https_proxy='http://127.0.0.1:9'))
check("fetch: a provider out of reach is said so, with how to read a copy a person downloaded",
      r.returncode == 1 and 'the provider did not answer' in r.stderr
      and 'python3 bin/fetch.py unece-rec20 --from <file>' in r.stderr, r.stderr)

# the round trip of what an invoice writes: its units, its calendar cells, a reading, and what crosses to nothing
U = os.path.join(ROOT, 'bin', 'units.py')


def units(*a):
    p = subprocess.run([PY, U] + list(a), capture_output=True, text=True, encoding='utf-8', env=dict(os.environ))
    return p.returncode, p.stdout.strip()


cells = {c: dmunits.read('unece-rec20', '3', c) for c in ('DAY', 'WEE', 'MON', 'ANN')}
check("UNECE: an invoice's DAY, WEE, MON and ANN are a calendar's cells — an extent's `level` and `count` — never a "
      "duration, and are written back as themselves",
      cells == {'DAY': {'of': 'time', 'in': 'gregorian-civil', 'level': 'day', 'count': '3'},
                'WEE': {'of': 'time', 'in': 'iso-week', 'level': 'week', 'count': '3'},
                'MON': {'of': 'time', 'in': 'gregorian-civil', 'level': 'month', 'count': '3'},
                'ANN': {'of': 'time', 'in': 'gregorian-civil', 'level': 'year', 'count': '3'}}
      and all(dmunits.write('unece-rec20', v) == ('3', c) for c, v in cells.items())
      and units('3', 'unece-rec20:MON', 'd')[0] == 1 and units('1', 'd', 'unece-rec20')[0] == 1, cells)
check("UNECE: KGM, KWH and HUR are the law's kg, kW.h and h, converted in the law and written back as their codes; a "
      "piece and C62 are one count of things, written back as C62, the one marked `back`",
      units('12.5', 'unece-rec20:KGM') == (0, '12.5 kg') and units('1.5', 'unece-rec20:KWH', 'J') == (0, '5400000 J')
      and units('2', 'unece-rec20:HUR', 'min') == (0, '120 min') and units('12.5', 'kg', 'unece-rec20') ==
      (0, '12.5 unece-rec20:KGM') and units('5', 'unece-rec20:H87', 'unece-rec20') == (0, '5 unece-rec20:C62'),
      [units('12.5', 'unece-rec20:KGM'), units('1.5', 'unece-rec20:KWH', 'J')])
check("UNECE: a Celsius reading is a position on its scale, never a quantity; a kelvin reading below zero is no position",
      units('21.5', 'unece-rec20:CEL') == (0, 'C:21.5') and dmunits.write('unece-rec20', 'C:21.5') == ('21.5', 'CEL')
      and units('-3', 'unece-rec20:KEL')[0] == 1, units('21.5', 'unece-rec20:CEL'))
refusals = {c: units('1', f'unece-rec20:{c}') for c in ('J2', '4P', 'GL', 'DZN', 'ZZ', 'QQQ')}
check("UNECE: a code with no kind in the law crosses to nothing and says why — J2 and 4P waiting for a kind (vacant, "
      "universal); GL, DZN and ZZ out of the law's units; a code no row names refused",
      all(rc == 1 for rc, _o in refusals.values()) and 'vacant: universal' in refusals['J2'][1]
      and 'surface tension' in refusals['4P'][1] and 'density' in refusals['GL'][1] and 'twelve' in refusals['DZN'][1]
      and 'no row of its crossing names it' in refusals['QQQ'][1], refusals)

# ...and saved through the gate: what an invoice wrote, read in its codes, is a garden's statements
SHOP = tempfile.mkdtemp(prefix='core-standards-shop-')
G = os.path.join(SHOP, 'shop')
REL = grow.release(os.path.join(SHOP, 'release'))
r = grow.garden(REL, G, 'sam', '--gardener-name', 'Sam')
hire = dict(dmunits.read('unece-rec20', '3', 'MON'), **{'from': '2026-10-01'})
mass = dmunits.read('unece-rec20', '12.5', 'KGM')
with open(os.path.join(G, 'beans', 'lathe.md'), 'w', encoding='utf-8', newline='\n') as fh:
    fh.write('---\nbean: lathe\nkind: host\ntitle: "lathe — the lathe hired, weighed at delivery"\nstatements:\n'
             '  - read: { by: sam, at: now }\n  - own: { by: sam, of: self }\n'
             f'  - measure: {{ id: weight, by: sam, of: self, as: mass, value: {{ count: "{mass["count"]}", unit: '
             f'{mass["unit"]} }}, at: "2026-10-01 10:00+03:00" }}\n---\nThe lathe.\n')
with open(os.path.join(G, 'beans', 'lathe-hire.md'), 'w', encoding='utf-8', newline='\n') as fh:
    fh.write('---\nbean: lathe-hire\nkind: contract\ntitle: "lathe-hire — the lathe, hired for three months"\n'
             'statements:\n  - say: { by: sam, at: now }\n  - own: { by: theone, of: self }\n'
             '  - answer: { by: sam, of: self, as: law }\n  - agree: { by: sam, of: lathe, through: written, as: hirer }\n'
             '  - pay: { id: rent, by: sam, of: { count: "300.00", unit: XTS } }\n'
             '  - obligatory: { id: monthly, of: rent, through: self, note: "the invoice\'s line: 3 MON of hire", '
             f'clause: {{ during: {{ of: {hire["of"]}, in: {hire["in"]}, level: {hire["level"]}, count: "{hire["count"]}", '
             f'from: "{hire["from"]}" }} }} }}\n---\nThe hire.\n')
grow.run('git', 'add', '-A', cwd=G)
s = grow.run(PY, 'bin/save.py', 'sam', 'the lathe and its hire, from the invoice', '--body',
             '- action: wrote [[lathe]] and [[lathe-hire]], their units read in UNECE Recommendation 20', cwd=G)
check("UNECE: an invoice's 12.5 KGM and 3 MON, read in its codes, are a measure in kg and a clause's extent of three "
      "months — saved through the gate", r.returncode == 0 and s.returncode == 0, (r.out[-300:], s.out[-900:]))

# ANY CODE LIST CROSSES BY THE ONE FORM, a garden's own too: its gauges, read through what each reads and how it shows it
# (two axes, as LOINC's property and scale), and its weighings, a family whose last digit says the decimal places
VOCAB = ('---\nnamespaces: []\nschemes:\n'
         '  - { scheme: gauges, classifies: "the shop\'s gauges, by what each reads and how it shows it", holding: '
         'extract, publisher: the shop, url: "file:extracts/gauges.tsv", levels: [ { level: gauge } ], neighbours: none, '
         'sources: extracts/gauges.tsv, units: units-gauges }\n'
         '  - { scheme: weighings, classifies: "the shop\'s weighing slips, each code saying its unit and its places", '
         'holding: extract, publisher: the shop, url: "file:extracts/weighings.tsv", levels: [ { level: slip } ], '
         'neighbours: none, sources: extracts/weighings.tsv, units: units-weighings }\n'
         'files:\n  - { registry: gauges, file: extracts/gauges.tsv, key: code }\n'
         '  - { registry: units-gauges, file: extracts/units-gauges.tsv, key: code }\n'
         '  - { registry: weighings, file: extracts/weighings.tsv, key: code }\n'
         '  - { registry: units-weighings, file: extracts/units-weighings.tsv, key: code }\n'
         "---\n# the shop's own rows\n")
EXTRACTS = {
    'gauges.tsv': 'code\tname\treads\tshows\noven\tthe oven\'s thermometer\theat\tnumber\n'
                  'scale\tthe platform scale\tweight\tnumber\npress\tthe press\'s stage\tstage\trank\n'
                  'log\tthe press\'s log\tremark\twords\ndoor\tthe oven\'s door\tajar\twords\n',
    'units-gauges.tsv': 'axis\tcode\tquantity\tline\tshape\twhy\nreads\theat\t\ttemperature\t\t\n'
                        'reads\tweight\tmass\t\t\t\nreads\tstage\t\tordinal\t\t\nshows\tnumber\t\t\tquantity\t\n'
                        'shows\trank\t\t\tposition\t\nreads\tremark\t\t\t\ta remark: nothing measured\n'
                        'shows\twords\t\t\ttext\t\n',
    'weighings.tsv': 'code\tname\nWT0\tkilograms, whole\nWT2\tkilograms, two places\nWT3\tkilograms, three places\n',
    'units-weighings.tsv': 'code\tunit\tdecimals\tvacant\twhy\nWTn\tkg\tn\t\t\nTARE\t\t\tuniversal\t'
                           'a slip\'s tare, which the shop has no kind for yet\n'}
with open(os.path.join(G, 'VOCAB.md'), 'w', encoding='utf-8', newline='\n') as fh:
    fh.write(VOCAB)
os.makedirs(os.path.join(G, 'extracts'), exist_ok=True)
for name, body in EXTRACTS.items():
    with open(os.path.join(G, 'extracts', name), 'w', encoding='utf-8', newline='\n') as fh:
        fh.write(body)
grow.run('git', 'add', '-A', cwd=G)
s = grow.run(PY, 'bin/save.py', 'sam', 'RULE-CHANGE: the gauges and the weighing slips, crossed into the law',
             '--body', '- action: RULE-CHANGE — VOCAB.md takes two schemes of the shop\'s own, each crossed into the '
             "law's units and kinds by a registry in extracts/", '- ratified_by: sam', cwd=G)
JUDGE = ("import sys, json; sys.path.insert(0, 'bin'); import units as u\n"
         "out = {}\n"
         "for k, (s, c, v) in json.loads(sys.argv[1]).items():\n"
         "    out[k] = u.judge(s, c, v)\n"
         "out['back'] = list(u.code_of('gauges', {'quantity': 'mass'}, 'reads'))[0]\n"
         "try:\n    u.through('gauges', 'door')\nexcept ValueError as e:\n    out['door'] = str(e)\n"
         "rows = u._garden_law().std.knowledge.rows('gauges')\n"
         "out['against'] = u.against('gauges', rows)[3]\n"
         "print(json.dumps(out))\n")
cases = {'oven-reading': ['gauges', 'oven', 'C:21.5'], 'oven-difference': ['gauges', 'oven', {'count': '2.5', 'unit': 'K'}],
         'scale-kg': ['gauges', 'scale', {'count': '12.5', 'unit': 'kg'}], 'scale-newton': ['gauges', 'scale',
         {'count': '120', 'unit': 'N'}], 'press-stage': ['gauges', 'press', 'ordinal:3'],
         'press-text': ['gauges', 'press', 'the third'], 'log-text': ['gauges', 'log', 'the oven ran hot'],
         'log-number': ['gauges', 'log', {'count': '3', 'unit': '{item}'}]}
j = grow.run(PY, '-c', JUDGE, json.dumps(cases), cwd=G)
got = json.loads(j.stdout) if j.returncode == 0 and j.stdout.strip() else {}
code = grow.run(PY, 'bin/units.py', '--code', 'gauges:oven', cwd=G)
check("a garden's own code list crosses by the same form: its gauges read through two axes — what each reads (a kind, "
      "or a line) and how it shows it (a shape) — saved through the gate",
      s.returncode == 0 and code.returncode == 0 and 'reads heat: line: temperature' in code.out
      and 'shows number: shape: quantity' in code.out, (s.out[-800:], code.out))
check("...a value is judged by what its code says through them: the oven's 21.5 °C a reading on its scale, never a "
      "difference; the scale's kilograms a mass, never newtons; the press's stage an ordinal position, never prose — and "
      "the law's mass written back as the axis value `weight`",
      got.get('oven-reading') == '' and 'is a position' in got.get('oven-difference', '')
      and got.get('scale-kg') == '' and 'measures force' in got.get('scale-newton', '')
      and got.get('press-stage') == '' and 'is a position' in got.get('press-text', '') and got.get('back') == 'weight',
      (got, j.out[-600:]))
check("...an axis value crossing to nothing is passed over and the other axis still read (the log's remark is no "
      "measure, and its words are text); a value no row names is refused, and the table's unknown values are listed",
      got.get('log-text') == '' and 'is text' in got.get('log-number', '')
      and "writes 'ajar' on its axis reads, and no row of its crossing names it" in got.get('door', '')
      and got.get('against') == {'reads': ['ajar']}, got)
w = [grow.run(PY, 'bin/units.py', *a, cwd=G) for a in (('001250', 'weighings:WT3'), ('1.25', 'kg', 'weighings'),
                                                        ('12', 'weighings:WT0'), ('1.25', 'weighings:WT2'),
                                                        ('1', 'weighings:TARE'))]
check("...and a family whose last digit says the decimal places (`decimals: n`): 001250 in WT3 is 1.250 kg, 1.25 kg is "
      "written back as 125 in WT2; a count written with a point, and a code crossing to nothing, refused",
      [x.out.strip() for x in w[:3]] == ['1.250 kg', '125 weighings:WT2', '12 kg'] and w[3].returncode == 1
      and 'digits without a point' in w[3].out and w[4].returncode == 1 and 'vacant: universal' in w[4].out,
      [x.out for x in w])
with open(os.path.join(G, 'extracts', 'units-weighings.tsv'), 'a', encoding='utf-8', newline='\n') as fh:
    fh.write('LBn\tlb\tn\t\t\n')
g = grow.run(PY, 'bin/check.py', cwd=G)
check("...and a garden's crossing is held to the form by its gate: a unit the law has not is refused, naming it",
      g.returncode != 0 and "units-weighings LBn: 'lb' is no unit of the law" in g.out, g.out[-600:])
shutil.rmtree(CACHE, ignore_errors=True)
shutil.rmtree(SHOP, ignore_errors=True)

print(f"\ncore_standards: {len(FAILS)} failed")
sys.exit(1 if FAILS else 0)

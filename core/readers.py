"""readers — every reader of the law's one table (core/law/formats.yaml `readers`), in one place (2.13).

A READER turns a format's records into the records of a model where no mapping can say how: a provider's columns
checked, a path split into its folders and its name, an id minted by a convention. Two kinds read here:

  the records of a model   file-tree, copies, exiftool: Records in, Records out, each with its `model` — what
                           bin/import.py runs a crosswalk over (`IMPORTED`)
  a provider's table       each a scheme the law holds `fetched` is read by (registries.yaml `knowledge_schemes`,
                           `fetch.reader`): the provider's file in, the rows of the scheme's file out (`FETCHED`,
                           bin/fetch.py)

A profile's readers live with its asset and are named by the table with their `home` (the bank statements'
adapters, assets/accounting/lib/statements.py). Each refuses what it cannot read by name, never guessing.
"""
import csv
import io
import json
import os
import re
import shutil

from core.formats import Record, Refused, escape, folder_id, content_id, local_moment, _path_bytes  # noqa: F401


# ================================================================================================ a model's records
def file_tree(records, volume, carrier=None, **_):
    """A listing of paths (path, type, size, hash, and where it says them ino and nlink) as the records of the volume,
    its folders, each folder's path beneath the volume, its files and its contents (formats.yaml, reader `file-tree`): a
    file of an inode of more than one name a `link`, held through the volume's one holding of the inode, an `inode`,
    given once. Each content once, after every file that holds it."""
    if carrier:
        yield Record({'id': volume, 'carrier': carrier}, at='the volume', model='volume')
    contents, inodes = {}, set()
    for rec in records:
        typ = rec.get('type')
        if typ not in ('d', 'f'):
            continue
        path = _path_bytes(rec.get('path', b''))
        if path in (b'', b'.'):
            continue
        parent, _, name = path.rpartition(b'/')
        pid = folder_id(volume, parent) if parent else volume
        if typ == 'd':
            fid = folder_id(volume, path)
            yield Record({'id': fid, 'parent': pid, 'name': name}, at=rec.at, model='folder')
            yield Record({'folder': fid, 'volume': volume, 'path': path}, at=rec.at, model='path')
            continue
        h = rec.get('hash')
        c = content_id(h) if h else 'unknown'
        if h and c not in contents:
            contents[c] = (str(h), rec.get('size'))
        ino, nlink = rec.get('ino'), rec.get('nlink')
        if h and ino not in (None, '') and str(nlink or '').isdigit() and int(nlink) > 1:
            iid = f"i{int(ino)}"
            if iid not in inodes:
                inodes.add(iid)
                yield Record({'volume': volume, 'id': iid, 'content': c}, at=rec.at, model='inode')
            yield Record({'folder': pid, 'content': c, 'name': name, 'through': f"{volume}#{iid}"}, at=rec.at,
                         model='link')
            continue
        yield Record({'folder': pid, 'content': c, 'name': name}, at=rec.at, model='file')
    for c, (h, size) in contents.items():
        yield Record({'id': c, 'blake3': h, 'size': '' if size is None else str(size)}, at=f"content {c}",
                     model='content')


def _tidy(path):
    """A path as a tree names it: no `.` folder and no empty one (`Edited/./a.mp4` is `Edited/a.mp4`); `..` is kept."""
    return b'/'.join(p for p in path.split(b'/') if p not in (b'', b'.'))


def copies(records, volume, known=None, ref=None, carrier=None, **_):
    """A record of copies (archive.py's manifest) as the folders of each destination not met before and each copy, its
    source the holding it was copied from (formats.yaml, reader `copies`), named by `ref` (verb, being, roles) as the
    keeping names a statement. `known` is the set of folders met already, kept between calls (a copier's one record at a
    time)."""
    known = set() if known is None else known
    if carrier and ('volume', volume) not in known:
        known.add(('volume', volume))
        yield Record({'id': volume, 'carrier': carrier}, at='the volume', model='volume')
    for rec in records:
        dest = _tidy(_path_bytes(rec.get('dest', b'')))
        parent, _, name = dest.rpartition(b'/')
        chain, p = [], parent
        while p and p not in known:
            chain.append(p)
            p = p.rpartition(b'/')[0]
        for p in reversed(chain):
            up, _, nm = p.rpartition(b'/')
            known.add(p)
            yield Record({'id': folder_id(volume, p), 'parent': folder_id(volume, up) if up else volume, 'name': nm},
                         at=rec.at, model='folder')
            yield Record({'folder': folder_id(volume, p), 'volume': volume, 'path': p}, at=rec.at, model='path')
        holder = folder_id(volume, parent) if parent else volume
        c = content_id(rec.get('h'))
        if ('content', c) not in known:                  # the content it holds, by its handle: a being of its kind
            known.add(('content', c))
            yield Record({'id': c}, at=rec.at, model='content')
        src_vol, src_path = str(rec.get('src_vol') or ''), _tidy(_path_bytes(rec.get('src_path', b'')))
        sp, _, sn = src_path.rpartition(b'/')
        sfolder = folder_id(src_vol, sp) if sp else src_vol
        source = ref('hold', sfolder, {'by': sfolder, 'of': c, 'as': escape(sn)}) if ref else None
        yield Record({'folder': holder, 'content': c, 'name': name, 'source': source,
                      'at': local_moment(rec.get('at')), 'note': f"archive.py copied it as {rec.get('class')}; its day "
                                                   f"{rec.get('day')} ({rec.get('day_from')})"},
                     at=rec.at, model='copy')


def exiftool(records, volume, root, find=None, model='image', **_):
    """exiftool's records, each file's tags and the content its place holds (formats.yaml, reader `exiftool`): `find`
    (volume, folder id, name) -> the content held there, or None. Each record is of the crosswalk's model (`image`, or
    `movie` for a video's header); `SourceFile`, which it is found by, is the reader's own and no field of it."""
    root = os.path.abspath(root)
    for rec in records:
        src = str(rec.get('SourceFile', ''))
        rel = os.path.relpath(os.path.abspath(src), root)
        out = Record({k: v for k, v in rec.items() if k != 'SourceFile'}, at=rec.at, model=model)
        if rel.startswith('..'):
            out['SourceFile'] = src                     # outside the root: kept, and listed as not carried
            yield out
            continue
        parent, _, name = os.fsencode(rel).rpartition(b'/')
        holder = folder_id(volume, parent) if parent else volume
        c = find(holder, escape(name)) if find else None
        if c:
            out['content'] = c
        yield out


IMPORTED = {'file-tree': file_tree, 'copies': copies, 'exiftool': exiftool}


def file_tree_sql(db, src, volume, carrier=None, **_):
    """The `file-tree` reader set-wise, inside the store's own engine (bin/import.py): the records of each model as a
    temporary table of the store's connection, read by SQL from the listing attached as `src` — the same records
    `file_tree` gives one at a time, the same ids. Returns ({model: table}, {column not read: how many rows hold it})."""
    import hashlib

    def fid(p):
        p = bytes(p) if p is not None else b''
        return 'd-' + hashlib.sha256(volume.encode() + b'\0' + p).hexdigest()[:16] if p else volume
    db.create_function('ft_fid', 1, fid, deterministic=True)
    db.create_function('ft_parent', 1, lambda p: bytes(p).rpartition(b'/')[0], deterministic=True)
    db.create_function('ft_name', 1, lambda p: bytes(p).rpartition(b'/')[2], deterministic=True)
    db.create_function('ft_bytes', 1, lambda p: p if isinstance(p, bytes) else os.fsencode(str(p)), deterministic=True)
    cols = [r[1] for r in db.execute(f'PRAGMA {src}.table_info(entry)')]
    if not {'path', 'type'} <= set(cols):
        raise Refused("the listing holds no table `entry` of paths and their types")
    has = lambda c: c if c in cols else 'NULL'
    for t in ('ft_volume', 'ft_folder', 'ft_path', 'ft_file', 'ft_link', 'ft_inode', 'ft_content_now'):
        db.execute(f'DROP TABLE IF EXISTS temp.{t}')
    db.execute("CREATE TABLE IF NOT EXISTS temp.ft_content (id TEXT PRIMARY KEY, blake3 TEXT, size TEXT)")
    db.execute("CREATE TEMP TABLE ft_volume (id TEXT, carrier TEXT)")
    if carrier:
        db.execute("INSERT INTO temp.ft_volume VALUES (?, ?)", (volume, carrier))
    sel = (f"SELECT ft_bytes(path) AS p, type, {has('size')} AS size, {has('hash')} AS hash, {has('ino')} AS ino, "
           f"{has('nlink')} AS nlink FROM {src}.entry WHERE type IN ('d', 'f') AND path NOT IN (x'', '.', x'2e')")
    linked = "(hash IS NOT NULL AND hash != '' AND ino IS NOT NULL AND CAST(nlink AS INTEGER) > 1)"
    db.execute(f"CREATE TEMP TABLE ft_folder AS SELECT ft_fid(p) AS id, ft_fid(ft_parent(p)) AS parent, ft_name(p) AS name "
               f"FROM ({sel}) WHERE type = 'd'")
    db.execute(f"CREATE TEMP TABLE ft_path AS SELECT ft_fid(p) AS folder, ? AS volume, p AS path FROM ({sel}) "
               f"WHERE type = 'd'", (volume,))
    db.execute(f"CREATE TEMP TABLE ft_file AS SELECT ft_fid(ft_parent(p)) AS folder, CASE WHEN hash IS NULL OR hash = '' "
               f"THEN 'unknown' ELSE 'c-' || substr(hash, 1, 24) END AS content, ft_name(p) AS name, hash, size "
               f"FROM ({sel}) WHERE type = 'f' AND NOT {linked}")
    db.execute(f"CREATE TEMP TABLE ft_link AS SELECT ft_fid(ft_parent(p)) AS folder, 'c-' || substr(hash, 1, 24) AS content, "
               f"ft_name(p) AS name, ? || '#i' || CAST(ino AS INTEGER) AS through, hash, size FROM ({sel}) "
               f"WHERE type = 'f' AND {linked}", (volume,))
    db.execute(f"CREATE TEMP TABLE ft_inode AS SELECT ? AS volume, 'i' || CAST(ino AS INTEGER) AS id, "
               f"min('c-' || substr(hash, 1, 24)) AS content FROM ({sel}) WHERE type = 'f' AND {linked} "
               f"GROUP BY CAST(ino AS INTEGER)", (volume,))
    for t in ('ft_file', 'ft_link'):
        db.execute(f"INSERT OR IGNORE INTO temp.ft_content SELECT 'c-' || substr(hash, 1, 24), hash, "
                   f"coalesce(CAST(size AS TEXT), '') FROM temp.{t} WHERE hash IS NOT NULL AND hash != '' ORDER BY rowid")
    missed = {}
    for c in cols:
        if c not in ('path', 'type', 'size', 'hash', 'ino', 'nlink'):
            n = db.execute(f'SELECT COUNT("{c}") FROM {src}.entry').fetchone()[0]
            if n:
                missed[c] = n
    return {'volume': 'temp.ft_volume', 'folder': 'temp.ft_folder', 'path': 'temp.ft_path', 'file': 'temp.ft_file',
            'inode': 'temp.ft_inode',
            'link': 'temp.ft_link', 'content': 'temp.ft_content'}, missed


SETWISE = {'file-tree': file_tree_sql}


# ================================================================================================ a provider's table
def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise Refused("the provider's file is a spreadsheet (xlsx), and openpyxl, which reads one, is not installed "
                      "here (`pip install openpyxl`)") from None


def ilo_isco_08_xlsx(data):
    """The ILO's ISCO-08 structure (`ISCO-08 EN Structure and definitions.xlsx`): its first sheet, one row a group —
    `Level` (1 to 4), `ISCO 08 Code`, `Title EN` — read into code, level, parent and name_en. The definitions, tasks and
    occupations it also holds are the ILO's text, and are not read."""
    wb = _openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    head = [str(c).strip() if c is not None else '' for c in next(rows, ())]
    want = ('Level', 'ISCO 08 Code', 'Title EN')
    if any(w not in head for w in want):
        raise Refused(f"the ILO's file has not the columns read here ({', '.join(want)}): its first row is "
                      f"{', '.join(h for h in head if h)[:200]}")
    at = {w: head.index(w) for w in want}
    levels = {'1': 'major', '2': 'sub-major', '3': 'minor', '4': 'unit'}
    out = []
    for r in rows:
        if not r or all(c is None for c in r):
            continue
        level, code, title = (str(r[at[w]]).strip() if r[at[w]] is not None else '' for w in want)
        level = level[:-2] if level.endswith('.0') else level
        if level not in levels or not code.isdigit() or len(code) != int(level) or not title:
            raise Refused(f"a row of the ILO's file is no group read here: level {level!r}, code {code!r}")
        out.append({'code': code, 'level': levels[level], 'parent': code[:-1], 'name_en': title})
    if not out:
        raise Refused("the ILO's file holds no group")
    return out


def _head(cells):
    """A header row's cells as words: a cell's line breaks and the spreadsheet's escaped carriage returns are spaces."""
    return [' '.join(str(c).replace('_x000D_', ' ').split()) if c is not None else '' for c in cells]


# UN/CEFACT's marks in a code's `Status`, as Recommendation 20 has used them: `X` a code marked deleted, `D` deprecated.
# The others (an entry added, renamed, changed) say what moved in the revision, not whether the code stands.
REC20_STATUS = {'X': 'deleted', 'D': 'deprecated'}


def unece_rec20_xlsx(data):
    """UN/CEFACT's Recommendation No. 20 (`rec20_Rev<N>e-<year>.xlsx`): its sheet of every code, one row a code
    (Annexes II and III) — `Status`, `Common Code`, `Name`, `Symbol` — read into code, status, name_en and symbol, the
    status `deleted` (`X`) or `deprecated` (`D`), another mark as written, and none where the code stands. Annex I, which
    lists a code once for each quantity it measures (its `Quantity`), is not that sheet. The descriptions, levels and
    conversion factors the file also holds are UNECE's text, and are not read."""
    wb = _openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    want = ('Status', 'Common Code', 'Name', 'Symbol')
    sheets = []
    for sheet in wb.worksheets:
        rows = sheet.iter_rows(values_only=True)
        head = _head(next(rows, ()))
        if all(w in head for w in want) and 'Quantity' not in head:
            sheets.append((sheet, rows, head))
    if len(sheets) != 1:
        raise Refused(f"{len(sheets)} sheets of UNECE's file list every code once with the columns read here "
                      f"({', '.join(want)}), and one is read: its sheets are {', '.join(s.title for s in wb.worksheets)}")
    sheet, rows, head = sheets[0]
    at = {w: head.index(w) for w in want}
    out, seen = [], set()
    for r in rows:
        if not r or all(c is None for c in r):
            continue
        status, code, name, symbol = (' '.join(str(r[at[w]]).split()) if at[w] < len(r) and r[at[w]] is not None
                                      else '' for w in want)
        if not re.fullmatch(r'[0-9A-Z]{2,3}', code) or not name or code in seen:
            raise Refused(f"a row of UNECE's file is no code read here: code {code!r}, name {name[:60]!r}"
                          + (" (a second row of it)" if code in seen else ''))
        seen.add(code)
        out.append({'code': code, 'status': REC20_STATUS.get(status, status), 'name_en': name, 'symbol': symbol})
    if not out:
        raise Refused("UNECE's file holds no code")
    return out


def loinc_table_csv(data):
    """LOINC's table (`LoincTable/Loinc.csv`, alone or inside the release's zip): one row a term — `LOINC_NUM`, `PROPERTY`,
    `SCALE_TYP`, `STATUS`, `LONG_COMMON_NAME` — read into code, status, property, scale and name_en, the status as LOINC
    writes it in lower case (`active`, `trial`, `discouraged`, `deprecated`). The table's other columns — the component,
    the system, the method, the definitions, the example units — are LOINC's text, and are not read."""
    if data[:2] == b'PK':
        import zipfile
        try:
            z = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile as e:
            raise Refused(f"LOINC's file is no zip it can be read in: {e}") from None
        inside = [n for n in z.namelist() if n.replace('\\', '/').endswith('LoincTable/Loinc.csv')]
        if len(inside) != 1:
            raise Refused(f"LOINC's zip holds {len(inside)} LoincTable/Loinc.csv, and one is read")
        data = z.read(inside[0])
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise Refused("LOINC's table is not UTF-8 text: no LoincTable/Loinc.csv") from None
    rows = csv.reader(io.StringIO(text))
    head = [h.strip() for h in next(rows, [])]
    want = ('LOINC_NUM', 'PROPERTY', 'SCALE_TYP', 'STATUS', 'LONG_COMMON_NAME')
    if any(w not in head for w in want):
        raise Refused(f"LOINC's file has not the columns read here ({', '.join(want)}): its first row is "
                      f"{', '.join(h for h in head if h)[:200]} — the table is LoincTable/Loinc.csv of LOINC's release")
    at = {w: head.index(w) for w in want}
    out, seen = [], set()
    for r in rows:
        if not r or all(not c.strip() for c in r):
            continue
        code, prop, scale, status, name = (r[at[w]].strip() if at[w] < len(r) else '' for w in want)
        if not re.fullmatch(r'[1-9][0-9]{0,6}-[0-9]', code) or code in seen or not prop or not scale:
            raise Refused(f"a row of LOINC's table is no term read here: code {code!r}, property {prop!r}, scale {scale!r}"
                          + (" (a second row of it)" if code in seen else ''))
        seen.add(code)
        out.append({'code': code, 'status': status.lower(), 'property': prop, 'scale': scale, 'name_en': name})
    if not out:
        raise Refused("LOINC's table holds no term")
    return out


def ifc_express(data):
    """IFC's schema in EXPRESS (`IFC4X3_ADD2.exp`): its measure types (`TYPE Ifc…Measure = …;`) and its SI units — each
    name of `IfcSIUnitName`, alone and after each prefix of `IfcSIPrefix`, written as a model writes the two
    (`MILLI.METRE`) — read into code, kind (`measure` or `unit`) and base (the type a measure is defined by; a unit's
    name). The schema's entities, rules and functions are buildingSMART's text, and are not read."""
    try:
        text = data.decode('utf-8')
    except UnicodeDecodeError:
        raise Refused("IFC's schema is not UTF-8 text: no EXPRESS schema") from None
    if not re.search(r'^SCHEMA IFC\w*;', text, re.M):
        raise Refused("the file is no IFC schema in EXPRESS: it declares no `SCHEMA IFC…;`")

    def enum(name):
        m = re.search(r'^TYPE ' + name + r' = ENUMERATION OF\s*\(([^)]*)\);', text, re.M)
        return [v.strip() for v in m.group(1).split(',') if v.strip()] if m else []
    measures = re.findall(r'^TYPE (Ifc\w+Measure) = ([^;]+);', text, re.M)
    names, prefixes = enum('IfcSIUnitName'), enum('IfcSIPrefix')
    if not measures or not names or not prefixes:
        raise Refused(f"IFC's schema holds {len(measures)} measure types, {len(names)} SI unit names and {len(prefixes)} "
                      f"prefixes, and each is read")
    out = [{'code': m, 'kind': 'measure', 'base': ' '.join(b.split())} for m, b in measures]
    for n in names:
        out.append({'code': n, 'kind': 'unit', 'base': n})
        out += [{'code': f"{p}.{n}", 'kind': 'unit', 'base': n} for p in prefixes]
    return out


def _sql_inserts(text, table):
    """Each `INSERT INTO <table> (columns) VALUES (...);` of an SQL script, as {column: value}: a string with its doubled
    quotes undone, a number as written, NULL as ''. A statement may span lines, as a remark's line break makes it."""
    head = re.compile(r'INSERT INTO ' + re.escape(table) + r' \(([^)]*)\) VALUES \(')
    for m in head.finditer(text):
        cols, vals, i = [c.strip() for c in m.group(1).split(',')], [], m.end()
        while i < len(text) and text[i] != ')':
            if text[i] == "'":
                j, buf = i + 1, []
                while j < len(text) and not (text[j] == "'" and text[j + 1:j + 2] != "'"):
                    buf.append(text[j])
                    j += 2 if text[j] == "'" else 1
                vals.append(''.join(buf))
                i = j + 1
            elif text[i] in ' ,\t\r\n':
                i += 1
            else:
                j = i
                while j < len(text) and text[j] not in ',)':
                    j += 1
                v = text[i:j].strip()
                vals.append('' if v.upper() == 'NULL' else v)
                i = j
        if len(vals) != len(cols):
            raise Refused(f"an insert into {table} holds {len(vals)} values for {len(cols)} columns")
        yield dict(zip(cols, vals))


def epsg_sql(data):
    """EPSG's dataset as its SQL scripts (`EPSG-v<N>-PostgreSQL.zip`, or its `…Data_Script.sql`): each row of
    `epsg_unitofmeasure` — `uom_code`, `unit_of_meas_name`, `unit_of_meas_type`, `target_uom_code`, `factor_b`,
    `factor_c`, `deprecated` — read into code, name_en, type, target, factor_b, factor_c and status (`deprecated`, or
    none): a unit is factor_b / factor_c of its target, written as EPSG writes them. The dataset's other tables (the
    reference systems, the transformations) and its remarks are EPSG's, and are not read here."""
    if data[:2] == b'PK':
        import zipfile
        try:
            z = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile as e:
            raise Refused(f"EPSG's file is no zip it can be read in: {e}") from None
        inside = [n for n in z.namelist() if n.endswith('Data_Script.sql')]
        if len(inside) != 1:
            raise Refused(f"EPSG's zip holds {len(inside)} data scripts (…Data_Script.sql), and one is read")
        data = z.read(inside[0])
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise Refused("EPSG's data script is not UTF-8 text") from None
    want = ('uom_code', 'unit_of_meas_name', 'unit_of_meas_type', 'target_uom_code', 'factor_b', 'factor_c', 'deprecated')
    out, seen = [], set()
    for r in _sql_inserts(text, 'epsg_unitofmeasure'):
        if any(w not in r for w in want):
            raise Refused(f"EPSG's units of measure have not the columns read here ({', '.join(want)})")
        code = r['uom_code']
        if not code.isdigit() or code in seen or not r['unit_of_meas_name']:
            raise Refused(f"a unit of EPSG's dataset is none read here: code {code!r}"
                          + (" (a second row of it)" if code in seen else ''))
        seen.add(code)
        out.append({'code': code, 'name_en': r['unit_of_meas_name'], 'type': r['unit_of_meas_type'],
                    'target': r['target_uom_code'], 'factor_b': r['factor_b'], 'factor_c': r['factor_c'],
                    'status': 'deprecated' if r['deprecated'] == '1' else ''})
    if not out:
        raise Refused("EPSG's file holds no unit of measure (`INSERT INTO epsg_unitofmeasure`): its data script is read")
    return out


def gs1_syntax_dictionary(data):
    """GS1's Barcode Syntax Dictionary (`gs1-syntax-dictionary.txt`): one line an application identifier or a range of
    them — `AIs [flags] specification [attributes] # title` — read into code (each AI of a range: `3100-3105` is six),
    format (its specification: `N6`, `N3,iso4217 N..15`) and name_en (its title, `NET WEIGHT (kg)`). Its flags, its
    attributes and its comments are GS1's, and are not read."""
    try:
        text = data.decode('utf-8')
    except UnicodeDecodeError:
        raise Refused("GS1's dictionary is not UTF-8 text") from None
    out, seen = [], set()
    for n, line in enumerate(text.splitlines(), 1):
        if not line.strip() or line.lstrip().startswith('#'):
            continue
        body, _, title = line.partition('#')
        words = body.split()
        m = re.fullmatch(r'([0-9]{2,4})(?:-([0-9]{2,4}))?', words[0])
        if not m or (m.group(2) and (len(m.group(2)) != len(m.group(1)) or m.group(2) < m.group(1))):
            raise Refused(f"line {n} of GS1's dictionary names no application identifier read here: {words[0]!r}")
        spec = [w for w in words[1:] if re.match(r'^\[?[NXYZ][0-9.]', w)]
        if not spec:
            raise Refused(f"line {n} of GS1's dictionary gives {words[0]} no specification")
        first, last = m.group(1), m.group(2) or m.group(1)
        for ai in range(int(first), int(last) + 1):
            code = str(ai).zfill(len(first))
            if code in seen:
                raise Refused(f"GS1's dictionary names the application identifier {code} twice")
            seen.add(code)
            out.append({'code': code, 'format': ' '.join(spec), 'name_en': ' '.join(title.split())})
    if not out:
        raise Refused("GS1's dictionary names no application identifier")
    return out


CLDR_LANGUAGES = ('en', 'fa', 'tr')     # the names read: the law's English, and the languages of its widgets


def cldr_units_json(data):
    """Unicode CLDR's unit names (`cldr-units-full`, the package Unicode publishes, a gzipped tar; or a zip of
    cldr-json): for each unit of `main/en/units.json` — each entry of its `long` names that has a `displayName`, its
    prefixes and patterns passed over — the code, and its long display name in English, Persian and Turkish, read into
    code, name_en, name_fa and name_tr. Its patterns (`{0} metres`), short and narrow forms and other languages are
    CLDR's, and are not read."""
    import tarfile
    import zipfile
    files = {}
    try:
        if data[:2] == b'PK':
            z = zipfile.ZipFile(io.BytesIO(data))
            names = z.namelist()
            get = z.read
        else:
            tf = tarfile.open(fileobj=io.BytesIO(data), mode='r:*')
            names = [m.name for m in tf.getmembers() if m.isfile()]
            get = lambda n: tf.extractfile(n).read()  # noqa: E731
    except (zipfile.BadZipFile, tarfile.TarError) as e:
        raise Refused(f"CLDR's file is neither a tar nor a zip it can be read in: {e}") from None
    for lang in CLDR_LANGUAGES:
        hit = [n for n in names if n.replace('\\', '/').endswith(f'main/{lang}/units.json')]
        if len(hit) != 1:
            raise Refused(f"CLDR's file holds {len(hit)} main/{lang}/units.json, and one of each of "
                          f"{', '.join(CLDR_LANGUAGES)} is read")
        try:
            files[lang] = json.loads(get(hit[0]).decode('utf-8'))['main'][lang]['units']['long']
        except (ValueError, KeyError) as e:
            raise Refused(f"CLDR's main/{lang}/units.json holds no units' long names: {e}") from None
    out = []
    for code, entry in files['en'].items():
        if not (isinstance(entry, dict) and entry.get('displayName')):
            continue
        out.append({'code': code, **{f'name_{lang}': str((files[lang].get(code) or {}).get('displayName') or '')
                                     for lang in CLDR_LANGUAGES}})
    if not out:
        raise Refused("CLDR's English units name no unit")
    return out


# GeoNames' table of a country (its readme.txt, "the main 'geoname' table"): nineteen columns, tab-separated, UTF-8.
GEONAMES_COLUMNS = ('geonameid', 'name', 'asciiname', 'alternatenames', 'latitude', 'longitude', 'feature class',
                    'feature code', 'country code', 'cc2', 'admin1 code', 'admin2 code', 'admin3 code', 'admin4 code',
                    'population', 'elevation', 'dem', 'timezone', 'modification date')
_DEGREES = re.compile(r'-?[0-9]{1,3}(\.[0-9]+)?')


def geonames_dump(data):
    """GeoNames' export of a country (`export/dump/<CC>.zip`, or the `<CC>.txt` inside it): one line a feature, in the
    nineteen columns its readme.txt names — read into code (the geonameid), name, ascii, names (its alternate names, as
    GeoNames writes them, comma-separated), latitude, longitude (WGS 84, degrees as written), class, feature (its
    feature class and code), country, admin1, admin2 (the divisions' codes), population, timezone and modified (the
    row's last change). Its cc2, admin3, admin4, elevation and dem are GeoNames' too, and are not read. The release is
    the day the zip was written."""
    release = None
    if data[:2] == b'PK':
        import zipfile
        try:
            z = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile as e:
            raise Refused(f"GeoNames' file is no zip it can be read in: {e}") from None
        inside = [i for i in z.infolist() if re.fullmatch(r'[A-Z]{2}\.txt|no-country\.txt', os.path.basename(i.filename))]
        if len(inside) != 1:
            raise Refused(f"GeoNames' zip holds {len(inside)} tables of a country (<CC>.txt), and one is read")
        release = 'export of %04d-%02d-%02d' % inside[0].date_time[:3]
        data = z.read(inside[0])
    try:
        text = data.decode('utf-8')
    except UnicodeDecodeError:
        raise Refused("GeoNames' table is not UTF-8 text") from None
    out, seen = [], set()
    for n, line in enumerate(text.split('\n'), 1):
        if not line.strip():
            continue
        cells = line.rstrip('\r').split('\t')
        if len(cells) != len(GEONAMES_COLUMNS):
            raise Refused(f"line {n} of GeoNames' table has {len(cells)} columns, and its readme names "
                          f"{len(GEONAMES_COLUMNS)}: it is no table of a country read here")
        r = dict(zip(GEONAMES_COLUMNS, cells))
        code = r['geonameid']
        if not re.fullmatch(r'[1-9][0-9]{0,9}', code) or code in seen or not r['name'] \
                or not _DEGREES.fullmatch(r['latitude']) or not _DEGREES.fullmatch(r['longitude']):
            raise Refused(f"line {n} of GeoNames' table is no feature read here: geonameid {code!r}, name "
                          f"{r['name'][:60]!r}, at {r['latitude']!r},{r['longitude']!r}"
                          + (" (a second row of it)" if code in seen else ''))
        seen.add(code)
        out.append({'code': code, 'name': r['name'], 'ascii': r['asciiname'], 'names': r['alternatenames'],
                    'latitude': r['latitude'], 'longitude': r['longitude'], 'class': r['feature class'],
                    'feature': r['feature code'], 'country': r['country code'], 'admin1': r['admin1 code'],
                    'admin2': r['admin2 code'], 'population': r['population'], 'timezone': r['timezone'],
                    'modified': r['modification date']})
    if not out:
        raise Refused("GeoNames' table holds no feature")
    return {'rows': out, 'release': release}


# Photon's layers of an address (a Place's `address_type`) that name where a thing is at the scale a person says it:
# the rest — a street, a house, an `other` — are what a place search finds, not a gazetteer's rows.
PHOTON_LAYERS = ('country', 'state', 'county', 'city', 'district', 'locality')
PHOTON_ELEMENTS = {'N': 'node', 'W': 'way', 'R': 'relation'}
PLACE_LANGUAGES = ('en', 'fa', 'tr')     # the names read beside a place's own: the law's English, its widgets' languages


def _zstd_lines(data):
    """The lines of a zstd-compressed file, as text: by Python's own zstd (3.14), else by the `zstd` program, streamed —
    a country's dump is gigabytes once decompressed, and is never held whole."""
    try:
        from compression import zstd
        yield from io.TextIOWrapper(zstd.ZstdFile(io.BytesIO(data)), encoding='utf-8')
        return
    except ImportError:
        pass
    exe = shutil.which('zstd')
    if not exe:
        raise Refused("the dump is compressed with zstd, which this Python cannot read and no `zstd` program here "
                      "can: decompress it (zstd -d <file>) and read the .jsonl, with --from")
    import subprocess
    import tempfile
    import threading
    with tempfile.TemporaryDirectory(prefix='daftar-fetch-') as tmp:
        src = os.path.join(tmp, 'dump.zst')
        with open(src, 'wb') as fh:
            fh.write(data)
        p = subprocess.Popen([exe, '-dcq', src], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        err = []
        t = threading.Thread(target=lambda: err.append(p.stderr.read()), daemon=True)
        t.start()
        try:
            yield from io.TextIOWrapper(p.stdout, encoding='utf-8')
        finally:
            p.stdout.close()
            code = p.wait()
            t.join()
        if code != 0:
            raise Refused(f"zstd could not decompress the dump: {(err[0] if err else b'').decode(errors='replace')[:200]}")


def photon_dump(data):
    """A country's dump in the format Photon reads (GraphHopper's `photon-dump-<country>-<version>-latest.jsonl.zst`,
    or the `.jsonl` decompressed): its first object the file's header (`NominatimDumpFile`), then the places — read,
    of each place at one of PHOTON_LAYERS with a name and a centroid, into code (its element, `relation/6663864`),
    layer, kind (its tag, `place=city`), name (its own), name_en, name_fa, name_tr, latitude, longitude (its centroid,
    WGS 84) and country. Its addresses, extra tags, importance, bounding box and geometry are the dump's, and are not
    read; an element the dump gives twice is read once. The release is the header's data timestamp."""
    out, seen, release = [], set(), None
    try:
        lines = _zstd_lines(data) if data[:4] == b'\x28\xb5\x2f\xfd' else iter(io.StringIO(data.decode('utf-8')))
        for n, line in enumerate(lines, 1):
            if not line.strip():
                continue
            try:
                o = json.loads(line)
            except ValueError:
                raise Refused(f"line {n} of the dump is no JSON object") from None
            if n == 1:
                head = o.get('content') if isinstance(o, dict) and o.get('type') == 'NominatimDumpFile' else None
                if not isinstance(head, dict) or not str(head.get('version', '')).startswith('0.1.'):
                    raise Refused("the file is no dump Photon reads: its first object is not a NominatimDumpFile of "
                                  "the format 0.1")
                release = f"data of {head.get('data_timestamp') or 'an unstated day'}, by {head.get('generator')} " \
                          f"{head.get('database_version') or ''}".strip()
                continue
            if not isinstance(o, dict) or o.get('type') != 'Place' or not isinstance(o.get('content'), list):
                continue
            for p in o['content']:
                if not isinstance(p, dict) or p.get('address_type') not in PHOTON_LAYERS:
                    continue
                names, c = p.get('name') or {}, p.get('centroid')
                kind, oid = PHOTON_ELEMENTS.get(p.get('object_type')), p.get('object_id')
                if not (kind and isinstance(oid, int) and oid > 0 and isinstance(names, dict) and names.get('name')
                        and isinstance(c, list) and len(c) == 2 and all(isinstance(x, (int, float)) for x in c)):
                    continue
                code = f"{kind}/{oid}"
                if code in seen:
                    continue
                seen.add(code)
                out.append({'code': code, 'layer': p['address_type'], 'kind': f"{p.get('osm_key')}={p.get('osm_value')}",
                            'name': ' '.join(str(names['name']).split()),
                            **{f'name_{l}': ' '.join(str(names.get(f'name:{l}') or '').split()) for l in PLACE_LANGUAGES},
                            'latitude': repr(float(c[1])), 'longitude': repr(float(c[0])),
                            'country': str(p.get('country_code') or '').upper()})
    except UnicodeDecodeError:
        raise Refused("the dump is not UTF-8 text") from None
    if release is None:
        raise Refused("the file is empty: no dump Photon reads")
    if not out:
        raise Refused("the dump holds no place at the scale of a country's divisions, cities and localities")
    return {'rows': out, 'release': release}


SMI = '{http://www.freedesktop.org/standards/shared-mime-info}'


def shared_mime_info(data):
    """freedesktop.org's shared-mime-info database: the release's tarball (its data/freedesktop.org.xml.in) or a host's
    installed copy (/usr/share/mime/packages/freedesktop.org.xml). Its document type declares the elements in an
    internal subset; one that declares an entity is refused, and the subset is set aside before the rest is read as
    XML, so nothing is expanded. Each top-level type once (code, level `top`, name_en), then each type: code, level
    `type`, parent its top-level type, name_en its comment in English, family its generic icon — or, where it gives
    none, its top-level type and `-x-generic`, as the specification makes one —, globs, sub_class_of and aliases, each
    a list written with spaces."""
    import tarfile
    import xml.etree.ElementTree as ET
    release = None
    if data[:2] == b'\x1f\x8b' or data[257:262] == b'ustar':
        try:
            tf = tarfile.open(fileobj=io.BytesIO(data), mode='r:*')
            names = [m.name for m in tf.getmembers() if m.isfile()]
        except tarfile.TarError as e:
            raise Refused(f"shared-mime-info's file is no tar it can be read in: {e}") from None
        hit = [n for n in names if re.search(r'(^|/)data/freedesktop\.org\.xml(\.in)?$', n)]
        if len(hit) != 1:
            raise Refused(f"the tarball holds {len(hit)} data/freedesktop.org.xml.in, and one is read")
        m = re.match(r'^shared-mime-info-([0-9][0-9.]*)/', hit[0])
        release = f"shared-mime-info {m.group(1)}" if m else None
        data = tf.extractfile(hit[0]).read()
    if re.search(rb'<!ENTITY', data):
        raise Refused("the database declares an entity: it is data, refused before it is read")
    body = re.sub(rb'<!DOCTYPE\s+mime-info\s*\[.*?\]\s*>', b'', data, count=1, flags=re.S)
    if re.search(rb'<!DOCTYPE', body):
        raise Refused("the database declares a document type other than shared-mime-info's own")
    try:
        root = ET.fromstring(body)
    except ET.ParseError as e:
        raise Refused(f"the database is not XML it can be read in: {e}") from None
    lang = '{http://www.w3.org/XML/1998/namespace}lang'
    tops, types = {}, []
    for t in root.iter(SMI + 'mime-type'):
        code = t.get('type') or ''
        if '/' not in code:
            continue
        top = code.split('/', 1)[0]
        comment = next((c.text for c in t.findall(SMI + 'comment') if c.get(lang) is None and c.text), '')
        icon = t.find(SMI + 'generic-icon')
        types.append({'code': code, 'level': 'type', 'parent': top, 'name_en': (comment or '').strip(),
                      'family': icon.get('name') if icon is not None else f"{top}-x-generic",
                      'globs': ' '.join(g.get('pattern') for g in t.findall(SMI + 'glob') if g.get('pattern')),
                      'sub_class_of': ' '.join(x.get('type') for x in t.findall(SMI + 'sub-class-of') if x.get('type')),
                      'aliases': ' '.join(x.get('type') for x in t.findall(SMI + 'alias') if x.get('type'))})
        tops.setdefault(top, {'code': top, 'level': 'top', 'parent': '', 'name_en': top, 'family': '', 'globs': '',
                              'sub_class_of': '', 'aliases': ''})
    if not types:
        raise Refused("the database names no media type")
    return {'rows': [tops[k] for k in sorted(tops)] + sorted(types, key=lambda r: r['code']), 'release': release}


# ================================================================================================ the beings' schemes
_OBO_TAG = re.compile(r'^([A-Za-z_-]+):\s?(.*)$')
_OBO_SYNONYM = re.compile(r'^"((?:[^"\\]|\\.)*)"\s+(EXACT|RELATED|NARROW|BROAD)\b')
_OBO_RANK = re.compile(r'^has_rank\s+NCBITaxon:([a-z_]+)')


def obo(data):
    """An ontology in the OBO flat file format (1.2 and 1.4), as the OBO Foundry's ontologies publish it — UBERON, the
    Cell Ontology, the Hymenoptera Anatomy and Plant ontologies, ChEBI, NCBITaxon — read one stanza at a time, never as
    a graph. The ontology's own terms alone (an id whose prefix is the header's `ontology`, its case aside; an imported
    term is its own ontology's): code (the id's local part, `UBERON:0000948` read `0000948`), level (the rank its
    `has_rank` names, NCBITaxon's; `term` where it names none), parent (its first `is_a` in the ontology), parents (each
    `is_a` there) and part_of (each `relationship: part_of` there), with spaces; name; synonyms (EXACT) and related
    (RELATED, NARROW, BROAD), each a list written with ` | `; status, `obsolete` where it is. Its definitions, comments
    and cross-references are the publisher's text, and are not read. The release is the header's `data-version`."""
    try:
        text = data.decode('utf-8')
    except UnicodeDecodeError:
        raise Refused("the ontology is not UTF-8 text") from None
    lines = text.split('\n')
    head, n = {}, 0
    for n, line in enumerate(lines):
        if line.startswith('['):
            break
        m = _OBO_TAG.match(line.rstrip('\r'))
        if m and m.group(1) not in head:
            head[m.group(1)] = m.group(2).strip()
    if 'format-version' not in head or not head.get('ontology'):
        raise Refused("the file is no OBO ontology: its header names no format-version and no ontology")
    prefix = head['ontology'].split('/')[-1].split('.')[0].upper() + ':'
    out, term, seen = [], None, set()

    def done(t):
        if not t or not t.get('id', '').upper().startswith(prefix):
            return
        code = t['id'][len(prefix):]
        if not code or code in seen:
            raise Refused(f"the ontology names the term {t['id']!r} " + ('twice' if code else 'with no local part'))
        seen.add(code)
        mine = [p[len(prefix):] for p in t['is_a'] if p.upper().startswith(prefix)]
        out.append({'code': code, 'level': t['rank'] or 'term', 'parent': mine[0] if mine else '',
                    'parents': ' '.join(mine), 'part_of': ' '.join(p[len(prefix):] for p in t['part_of']
                                                                   if p.upper().startswith(prefix)),
                    'name': t.get('name', ''), 'synonyms': ' | '.join(t['exact']), 'related': ' | '.join(t['related']),
                    'status': 'obsolete' if t['obsolete'] else ''})

    for line in lines[n:]:
        line = line.rstrip('\r')
        if line.startswith('['):
            done(term)
            term = {'is_a': [], 'part_of': [], 'exact': [], 'related': [], 'rank': '', 'obsolete': False} \
                if line.strip() == '[Term]' else None
            continue
        if term is None or not line or line.startswith('!'):
            continue
        m = _OBO_TAG.match(line)
        if not m:
            continue
        tag, value = m.group(1), m.group(2).split(' ! ')[0].strip()
        if tag == 'id':
            term['id'] = value
        elif tag == 'name':
            term['name'] = value
        elif tag == 'is_a':
            term['is_a'].append(value.split()[0])
        elif tag == 'relationship' and value.startswith('part_of '):
            term['part_of'].append(value.split()[1])
        elif tag == 'synonym':
            s = _OBO_SYNONYM.match(value)
            if s:
                (term['exact'] if s.group(2) == 'EXACT' else term['related']).append(
                    s.group(1).replace('\\"', '"').replace('|', '/'))
        elif tag == 'property_value':
            r = _OBO_RANK.match(value)
            if r:
                term['rank'] = r.group(1)
        elif tag == 'is_obsolete':
            term['obsolete'] = value == 'true'
    done(term)
    if not out:
        raise Refused(f"the ontology {head['ontology']!r} holds no term of its own ({prefix})")
    return {'rows': out, 'release': head.get('data-version') or None}


_GET_CODE = re.compile(r'^([A-Z]{1,3})([0-9]+)\.([0-9]+)$')


def iucn_get_maps(data):
    """The IUCN Global Ecosystem Typology's map details (`map-details.xml`, its maps of the ecosystem functional groups,
    published under CC BY 4.0): each functional group once — code (`T1.1`), level `group`, parent its biome (`T1`),
    name as the map names it after its code — and each biome and realm its codes name — `T1`, level `biome`, parent
    `T`; `T`, level `realm` — with no name: the biomes' and realms' names are in the typology's profiles, which IUCN
    publishes under its own copyright, and are not read. The release is the latest day a map was updated."""
    import xml.etree.ElementTree as ET
    if re.search(rb'<!(DOCTYPE|ENTITY)', data):
        raise Refused("the map details declare a document type or an entity: they are data, refused before they are read")
    try:
        root = ET.fromstring(data)
    except ET.ParseError as e:
        raise Refused(f"the map details are not XML they can be read in: {e}") from None
    groups, upper, days = {}, {}, []
    for m in root.iter('Map'):
        code = (m.get('efg_code') or '').strip()
        c = _GET_CODE.match(code)
        named = (m.findtext('Functional_group') or '').strip()
        if not c or not named.startswith(code + ' '):
            raise Refused(f"a map names the group {code!r} as {named[:60]!r}: no functional group read here")
        if m.get('update'):
            days.append(m.get('update'))
        realm, biome = c.group(1), c.group(1) + c.group(2)
        groups.setdefault(code, {'code': code, 'level': 'group', 'parent': biome, 'name': named[len(code):].strip()})
        upper.setdefault(biome, {'code': biome, 'level': 'biome', 'parent': realm, 'name': ''})
        upper.setdefault(realm, {'code': realm, 'level': 'realm', 'parent': '', 'name': ''})
    if not groups:
        raise Refused("the map details name no functional group")
    order = {'realm': 0, 'biome': 1, 'group': 2}
    rows = sorted(list(upper.values()) + list(groups.values()), key=lambda r: (order[r['level']], r['code']))
    return {'rows': rows, 'release': f"the maps updated to {max(days)}" if days else None}


COLDP_COLUMNS = ('col:ID', 'col:parentID', 'col:status', 'col:rank', 'col:scientificName', 'col:authorship')


def coldp(data):
    """A Catalogue of Life Data Package (ColDP), the zip ChecklistBank exports a release of the Catalogue of Life in:
    its NameUsage.tsv, one name usage a line — an accepted taxon, or a synonym whose parent is the taxon it names —
    read into code (`col:ID`), level (`col:rank`), parent (`col:parentID`), name (`col:scientificName`), authorship and
    status (`accepted`, `synonym`, …). Its metadata.yaml says the release: its `alias` (COL26.10 XR) and `issued`."""
    import zipfile
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise Refused(f"the package is no zip it can be read in: {e}") from None
    names = {os.path.basename(n): n for n in z.namelist()}
    if 'NameUsage.tsv' not in names:
        raise Refused("the package holds no NameUsage.tsv: a package of Name.tsv and Taxon.tsv is not read here")
    release = None
    if 'metadata.yaml' in names:
        meta = z.read(names['metadata.yaml']).decode('utf-8', 'replace')
        alias = re.search(r'^alias:\s*(.+)$', meta, re.M)
        issued = re.search(r'^issued:\s*(.+)$', meta, re.M)
        release = ' '.join(x.group(1).strip().strip('\'"') for x in (alias, issued) if x) or None
    with z.open(names['NameUsage.tsv']) as fh:
        lines = io.TextIOWrapper(fh, encoding='utf-8', newline='')
        head = lines.readline().rstrip('\r\n').split('\t')
        missing = [c for c in COLDP_COLUMNS if c not in head]
        if missing:
            raise Refused(f"NameUsage.tsv lacks the columns {', '.join(missing)}")
        at = [head.index(c) for c in COLDP_COLUMNS]
        out, seen = [], set()
        for n, line in enumerate(lines, 2):
            cells = line.rstrip('\r\n').split('\t')
            if len(cells) < len(head):
                if not line.strip():
                    continue
                raise Refused(f"line {n} of NameUsage.tsv has {len(cells)} cells, and its head names {len(head)}")
            code, parent, status, rank, name, auth = (cells[i] for i in at)
            if not code or code in seen or not name:
                raise Refused(f"line {n} of NameUsage.tsv is no name usage read here: id {code!r}, name {name[:60]!r}"
                              + (" (a second row of it)" if code in seen else ''))
            seen.add(code)
            out.append({'code': code, 'level': rank, 'parent': parent, 'name': name, 'authorship': auth,
                        'status': status})
    if not out:
        raise Refused("the package names no taxon")
    return {'rows': out, 'release': release}


def _etim_csv(z, name):
    try:
        raw = z.read(name)
    except KeyError:
        raise Refused(f"ETIM's export holds no {name}") from None
    text = raw.decode('utf-16' if raw[:2] in (b'\xff\xfe', b'\xfe\xff') else
                      'utf-16-le' if raw[1:2] == b'\x00' else 'utf-8-sig')
    return list(csv.DictReader(io.StringIO(text), delimiter=';'))


def etim_csv(data):
    """ETIM International's model in its CSV export (a zip of tables, UTF-16, `;` between cells): each group of
    ETIMARTGROUP.csv (code `EG…`, level `group`, name) and each class of ETIMARTCLASS.csv (code `EC…`, level `class`,
    parent its group, name, version), with the synonyms ETIMARTCLASSSYNONYMMAP.csv gives a class, written with ` | `.
    Its features, values and units are ETIM's too, and are not read."""
    import zipfile
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise Refused(f"ETIM's export is no zip it can be read in: {e}") from None
    groups = _etim_csv(z, 'ETIMARTGROUP.csv')
    classes = _etim_csv(z, 'ETIMARTCLASS.csv')
    synonyms = {}
    for r in _etim_csv(z, 'ETIMARTCLASSSYNONYMMAP.csv'):
        s = (r.get('CLASSSYNONYM') or '').strip()
        if s:
            synonyms.setdefault(r.get('ARTCLASSID'), []).append(s.replace('|', '/'))
    out = []
    for r in groups:
        if not re.fullmatch(r'EG[0-9]{6}', r.get('ARTGROUPID') or ''):
            raise Refused(f"ETIM's groups name {r.get('ARTGROUPID')!r}: no group read here")
        out.append({'code': r['ARTGROUPID'], 'level': 'group', 'parent': '', 'name': r.get('GROUPDESC') or '',
                    'synonyms': '', 'version': ''})
    known = {r['code'] for r in out}
    for r in classes:
        code, group = r.get('ARTCLASSID') or '', r.get('ARTGROUPID') or ''
        if not re.fullmatch(r'EC[0-9]{6}', code) or group not in known:
            raise Refused(f"ETIM's classes name {code!r} in the group {group!r}: no class read here")
        out.append({'code': code, 'level': 'class', 'parent': group, 'name': r.get('ARTCLASSDESC') or '',
                    'synonyms': ' | '.join(s for s in synonyms.get(code, []) if s != r.get('ARTCLASSDESC')),
                    'version': r.get('ARTCLASSVERSION') or ''})
    if len(out) == len(known):
        raise Refused("ETIM's export names no class")
    return out


_CPV_LEVELS = {2: 'division', 3: 'group', 4: 'class'}


def cpv_genericode(data):
    """The Common Procurement Vocabulary (CPV 2008) in the Genericode the Publications Office publishes it in: each row
    a code of eight digits (its check digit not written) — code, level by its digits not zero (two a `division`, three a
    `group`, four a `class`, five or more a `category`), parent the nearest code above it the list holds, name its
    English `Name`. Its other languages are the Publications Office's, and are not read. The release is the list's
    canonical version."""
    import xml.etree.ElementTree as ET
    if re.search(rb'<!(DOCTYPE|ENTITY)', data[:4096]):
        raise Refused("the code list declares a document type or an entity: it is data, refused before it is read")
    try:
        root = ET.fromstring(data)
    except ET.ParseError as e:
        raise Refused(f"the code list is not XML it can be read in: {e}") from None
    version = (root.findtext('Identification/CanonicalVersionUri') or '').rsplit('/', 1)[-1] or None
    names = {}
    for row in root.iter('Row'):
        cells = {v.get('ColumnRef'): (v.findtext('SimpleValue') or '').strip() for v in row.findall('Value')}
        code = cells.get('Code', '')
        if not re.fullmatch(r'[0-9]{8}', code) or code in names:
            raise Refused(f"the code list names {code!r}" + (' twice' if code in names else ': no code of the CPV'))
        names[code] = cells.get('Name') or cells.get('en_label') or ''
    if not names:
        raise Refused("the code list names no code")
    out = []
    for code in sorted(names):
        sig = len(code.rstrip('0')) or 1
        parent = ''
        for k in range(max(sig, 3) - 1, 1, -1):
            up = code[:k].ljust(8, '0')
            if up != code and up in names:
                parent = up
                break
        out.append({'code': code, 'level': _CPV_LEVELS.get(max(sig, 2), 'category'), 'parent': parent,
                    'name': names[code]})
    return {'rows': out, 'release': f"CPV 2008, the Publications Office's version {version}" if version else None}


FETCHED = {'ilo-isco-08-xlsx': ilo_isco_08_xlsx, 'unece-rec20-xlsx': unece_rec20_xlsx,
           'loinc-table-csv': loinc_table_csv, 'ifc-express': ifc_express, 'epsg-sql': epsg_sql,
           'gs1-syntax-dictionary': gs1_syntax_dictionary, 'cldr-units-json': cldr_units_json,
           'geonames-dump': geonames_dump, 'photon-dump': photon_dump, 'shared-mime-info': shared_mime_info,
           'obo': obo, 'coldp': coldp, 'etim-csv': etim_csv, 'cpv-genericode': cpv_genericode,
           'iucn-get-maps': iucn_get_maps}

#!/usr/bin/env python3
"""fetch — bring a standard's table from its provider to this machine: a scheme the law holds `fetched`
(core/law/registries.yaml, `knowledge_schemes`), downloaded from where its row's `fetch.from` says, into a cache outside
git, and read into the rows its `registry_files` file holds.

    python3 bin/fetch.py <scheme>                 # download it from its provider, and read it
    python3 bin/fetch.py <scheme> --from <file>   # a copy at hand: one this machine's person downloaded from the
                                                  # provider, or a peer's unchanged copy where the scheme's licence
                                                  # allows unchanged copies to be handed on
    python3 bin/fetch.py --list                   # what this machine has fetched, each copy by its SHA-256

DAFTAR SHIPS NO ONE ELSE'S TABLE where the provider's terms do not let it: each machine reaches the provider itself,
under the provider's terms, and keeps the provider's file unchanged beside the rows read from it. The cache is this
machine's, never the garden's: `DAFTAR_STANDARDS_CACHE`, else `%LOCALAPPDATA%\\daftar\\standards` on Windows, else
`$XDG_CACHE_HOME/daftar/standards` (`~/.cache/daftar/standards`). Each copy is kept under the SHA-256 of the provider's
file: `<scheme>/<sha-256>/source.<ext>`, `rows.tsv` and `fetched.json` (from where, when, how many rows). A garden reads
a fetched scheme's codes from the newest copy here (bin/knowledge.py), and from the release's copy while it carries one.

Nothing is sent but the request for the provider's file. Print the line it prints into the save's entry: what was
read, from where, and its SHA-256, so the garden says which copy its codes were read against.
"""
import csv
import datetime
import hashlib
import io
import json
import os
import re
import shutil
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import parse as dmparse  # noqa: E402,F401 — its import sets UTF-8 on stdout and stderr


class Refused(Exception):
    """Why a table was not fetched or not read: said to the person, and nothing is kept."""


def cache_dir():
    """Where this machine keeps what it fetched — outside every garden."""
    if os.environ.get('DAFTAR_STANDARDS_CACHE'):
        return os.environ['DAFTAR_STANDARDS_CACHE']
    if os.name == 'nt' and os.environ.get('LOCALAPPDATA'):
        return os.path.join(os.environ['LOCALAPPDATA'], 'daftar', 'standards')
    return os.path.join(os.environ.get('XDG_CACHE_HOME') or os.path.join(os.path.expanduser('~'), '.cache'),
                        'daftar', 'standards')


# ------------------------------------------------------------------------------------------------ the readers
def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise Refused("the provider's file is a spreadsheet (xlsx), and openpyxl, which reads one, is not installed "
                      "here (`pip install openpyxl`)") from None


def ilo_isco_08_xlsx(data):
    """The ILO's ISCO-08 structure (`ISCO-08 EN Structure and definitions.xlsx`): its first sheet, one row a group —
    `Level` (1 to 4), `ISCO 08 Code`, `Title EN` — read into code, level, parent and name_en. The definitions, tasks and
    occupations it also holds are the ILO's text, and are not read."""
    wb = _openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    head = [str(c).strip() if c is not None else '' for c in next(rows, ())]
    want = ('Level', 'ISCO 08 Code', 'Title EN')
    if any(w not in head for w in want):
        raise Refused(f"the ILO's file has not the columns read here ({', '.join(want)}): its first row is "
                      f"{', '.join(h for h in head if h)[:200]}")
    at = {w: head.index(w) for w in want}
    levels = {'1': 'major', '2': 'sub-major', '3': 'minor', '4': 'unit'}
    out = []
    for r in rows:
        if not r or all(c is None for c in r):
            continue
        level, code, title = (str(r[at[w]]).strip() if r[at[w]] is not None else '' for w in want)
        level = level[:-2] if level.endswith('.0') else level
        if level not in levels or not code.isdigit() or len(code) != int(level) or not title:
            raise Refused(f"a row of the ILO's file is no group read here: level {level!r}, code {code!r}")
        out.append({'code': code, 'level': levels[level], 'parent': code[:-1], 'name_en': title})
    if not out:
        raise Refused("the ILO's file holds no group")
    return out


def _head(cells):
    """A header row's cells as words: a cell's line breaks and the spreadsheet's escaped carriage returns are spaces."""
    return [' '.join(str(c).replace('_x000D_', ' ').split()) if c is not None else '' for c in cells]


# UN/CEFACT's marks in a code's `Status`, as Recommendation 20 has used them: `X` a code marked deleted, `D` deprecated.
# The others (an entry added, renamed, changed) say what moved in the revision, not whether the code stands.
REC20_STATUS = {'X': 'deleted', 'D': 'deprecated'}


def unece_rec20_xlsx(data):
    """UN/CEFACT's Recommendation No. 20 (`rec20_Rev<N>e-<year>.xlsx`): its sheet of every code, one row a code
    (Annexes II and III) — `Status`, `Common Code`, `Name`, `Symbol` — read into code, status, name_en and symbol, the
    status `deleted` (`X`) or `deprecated` (`D`), another mark as written, and none where the code stands. Annex I, which
    lists a code once for each quantity it measures (its `Quantity`), is not that sheet. The descriptions, levels and
    conversion factors the file also holds are UNECE's text, and are not read."""
    wb = _openpyxl().load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    want = ('Status', 'Common Code', 'Name', 'Symbol')
    sheets = []
    for sheet in wb.worksheets:
        rows = sheet.iter_rows(values_only=True)
        head = _head(next(rows, ()))
        if all(w in head for w in want) and 'Quantity' not in head:
            sheets.append((sheet, rows, head))
    if len(sheets) != 1:
        raise Refused(f"{len(sheets)} sheets of UNECE's file list every code once with the columns read here "
                      f"({', '.join(want)}), and one is read: its sheets are {', '.join(s.title for s in wb.worksheets)}")
    sheet, rows, head = sheets[0]
    at = {w: head.index(w) for w in want}
    out, seen = [], set()
    for r in rows:
        if not r or all(c is None for c in r):
            continue
        status, code, name, symbol = (' '.join(str(r[at[w]]).split()) if at[w] < len(r) and r[at[w]] is not None
                                      else '' for w in want)
        if not re.fullmatch(r'[0-9A-Z]{2,3}', code) or not name or code in seen:
            raise Refused(f"a row of UNECE's file is no code read here: code {code!r}, name {name[:60]!r}"
                          + (" (a second row of it)" if code in seen else ''))
        seen.add(code)
        out.append({'code': code, 'status': REC20_STATUS.get(status, status), 'name_en': name, 'symbol': symbol})
    if not out:
        raise Refused("UNECE's file holds no code")
    return out


READERS = {'ilo-isco-08-xlsx': ilo_isco_08_xlsx, 'unece-rec20-xlsx': unece_rec20_xlsx}


# ------------------------------------------------------------------------------------------------ the law's rows
def _law(root=ROOT):
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from core import read
    d = os.path.join(root, 'core', 'law')
    d = d if os.path.isfile(os.path.join(d, 'registries.yaml')) else os.path.join(ROOT, 'core', 'law')
    return read.data(os.path.join(d, 'registries.yaml'))


def scheme_row(scheme, root=ROOT):
    law = _law(root)
    row = next((r for r in law.get('knowledge_schemes') or [] if isinstance(r, dict) and r.get('scheme') == scheme), None)
    if row is None:
        raise Refused(f"{scheme!r} is no scheme of the knowledge tree")
    if row.get('holding') != 'fetched' or not isinstance(row.get('fetch'), dict):
        raise Refused(f"{scheme} is held `{row.get('holding') or 'shipped'}`, not fetched: its codes are read where "
                      f"the law holds them")
    columns = None
    f = next((r.get('file') for r in law.get('registry_files') or [] if isinstance(r, dict) and r.get('registry') == scheme),
             None)
    if f and os.path.isfile(os.path.join(ROOT, f)):
        with open(os.path.join(ROOT, f), encoding='utf-8') as fh:
            columns = next(csv.reader(fh, delimiter='\t'), None)
    return row, columns


# ------------------------------------------------------------------------------------------------ the act
def keep(scheme, data, source, *, root=ROOT, now=None):
    """Read the provider's file `data` (from `source`) by the scheme's reader, and keep it in this machine's cache.
    Returns (sha-256, rows, the directory kept in)."""
    row, columns = scheme_row(scheme, root)
    reader = READERS.get(str(row['fetch'].get('reader')))
    if reader is None:
        raise Refused(f"{scheme}'s reader `{row['fetch'].get('reader')}` is no reader of this release")
    rows = reader(data)
    if columns and set(columns) != set(rows[0]):
        raise Refused(f"the rows read have the columns {', '.join(rows[0])}, and {scheme}'s file has {', '.join(columns)}")
    sha = hashlib.sha256(data).hexdigest()
    d = os.path.join(cache_dir(), scheme, sha)
    tmp = d + '.part'
    shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp)
    ext = os.path.splitext(str(row['fetch'].get('from')).split('?')[0])[1] or '.bin'
    with open(os.path.join(tmp, 'source' + ext), 'wb') as fh:
        fh.write(data)
    cols = columns or list(rows[0])
    with open(os.path.join(tmp, 'rows.tsv'), 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=cols, delimiter='\t', lineterminator='\n')
        w.writeheader()
        w.writerows(rows)
    when = (now or datetime.datetime.now().astimezone()).isoformat(timespec='seconds')
    with open(os.path.join(tmp, 'fetched.json'), 'w', encoding='utf-8') as fh:
        json.dump({'scheme': scheme, 'from': source, 'at': when, 'sha256': sha, 'rows': len(rows)}, fh, indent=1)
    shutil.rmtree(d, ignore_errors=True)
    os.replace(tmp, d)
    return sha, rows, d


def download(url):
    import urllib.request
    req = urllib.request.Request(url, headers={'User-Agent': 'daftar-fetch (a standard\'s table, for its reader)'})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.read()
    except OSError as e:
        raise Refused(f"the provider did not answer ({url}): {e}") from None


def fetched(scheme):
    """The copies of `scheme` this machine keeps, newest first: [(when, sha-256, rows.tsv)]."""
    d = os.path.join(cache_dir(), scheme)
    out = []
    for sha in os.listdir(d) if os.path.isdir(d) else []:
        meta = os.path.join(d, sha, 'fetched.json')
        rows = os.path.join(d, sha, 'rows.tsv')
        if not sha.endswith('.part') and os.path.isfile(meta) and os.path.isfile(rows):
            try:
                with open(meta, encoding='utf-8') as fh:
                    m = json.load(fh)
            except (OSError, ValueError):
                continue
            out.append((str(m.get('at') or ''), sha, rows))
    return sorted(out, reverse=True)


def crossed(scheme, rows):
    """[line]: what the provider's rows, read, say of the scheme's crossing into the law's units and kinds (`units` on its
    row): how many of its codes cross (through every axis, where it names axes), each value on an axis no row names,
    and each crossed code it has not, or marks deprecated or deleted — a revision that moved what the law's crossing
    names. Nothing for a scheme with no crossing."""
    import units as dmunits
    try:
        n, missing, marked, unknown = dmunits.against(scheme, rows)
    except ValueError:
        return []
    out = [f"fetch: {n} of its {len(rows)} codes cross into the law's units and kinds (python3 bin/units.py --codes "
           f"{scheme})"]
    for axis, values in unknown.items():
        out.append(f"fetch: on its axis {axis}, values no row of the law's crossing names: {', '.join(values)}")
    if missing:
        out.append(f"fetch: crossed by the law, and not in this copy: {', '.join(missing)}")
    if marked:
        out.append(f"fetch: crossed by the law, and marked by the provider: "
                   f"{', '.join(f'{c} ({s})' for c, s in marked)}")
    return out


def main(argv):
    if not argv or argv[0] in ('-h', '--help'):
        print(__doc__)
        return 0 if argv else 2
    if argv == ['--list']:
        base = cache_dir()
        for scheme in sorted(os.listdir(base)) if os.path.isdir(base) else []:
            for when, sha, _rows in fetched(scheme):
                print(f"{scheme}\t{when}\tsha-256:{sha}")
        print(f"fetch: this machine's cache is {base}")
        return 0
    scheme, rest = argv[0], argv[1:]
    try:
        row, _cols = scheme_row(scheme)
        if rest[:1] == ['--from'] and len(rest) == 2:
            with open(rest[1], 'rb') as fh:
                data, source = fh.read(), os.path.abspath(rest[1])
        elif not rest:
            source = str(row['fetch'].get('from'))
            try:
                data = download(source)
            except Refused as e:
                raise Refused(f"{e} — download it from there by hand, as its provider serves a person, and read that "
                              f"copy: python3 bin/fetch.py {scheme} --from <file>") from None
        else:
            print(__doc__)
            return 2
        sha, rows, d = keep(scheme, data, source)
    except (Refused, OSError) as e:
        print(f"fetch: NOT FETCHED: {e}", file=sys.stderr)
        return 1
    print(f"fetch: {scheme}, {len(rows)} rows, kept in {d}")
    for line in crossed(scheme, rows):
        print(line)
    print(f"for the save's entry:\n  - read: {scheme} from {source}, sha-256 {sha}, {len(rows)} rows")
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
